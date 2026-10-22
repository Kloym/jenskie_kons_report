from __future__ import annotations

import hashlib
import queue
import re
import threading
from collections import Counter, defaultdict
from copy import copy
from dataclasses import dataclass, field
from datetime import date, timedelta
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO
from pathlib import Path
from calendar import monthrange

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.table import Table
from openpyxl.worksheet.views import Selection

from zhk_report import (
    AppConfig,
    HeaderResolver,
    ReportError,
    XlsWorkbookAdapter,
    clean_text,
    normalize_key,
    parse_date,
    service_code_from_cell,
    safe_excel_text,
)

UNITS = (
    "ЦЖЗ на Лобненской",
    "ЦЖЗ на Петрозаводской",
    "ЦЖЗ на Планетной",
    "ЖК №10",
)
PLANET = UNITS[2]

TEMPLATE = (
    Path(__file__).resolve().parent
    / "templates"
    / "analysis_14_template.xlsx"
)

MONTHS = (
    "",
    "Январь", "Февраль", "Март", "Апрель",
    "Май", "Июнь", "Июль", "Август",
    "Сентябрь", "Октябрь", "Ноябрь", "Декабрь",
)

ZERO = Decimal("0")
CENT = Decimal("0.01")


def number(value):
    text = clean_text(value).replace(" ", "").replace(",", ".")
    try:
        result = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError("пустое или неверное число") from exc

    if not result.is_finite():
        raise ValueError("число NaN/Infinity")

    return result


def code_key(value):
    text = safe_excel_text(value)
    if text.isdecimal():
        return text.zfill(6)
    return text.upper()


def identify_department(value):
    key = normalize_key(value)

    if not (
        "женск" in key
        or re.search(r"\b(?:жк|цжз)\b", key)
    ):
        return None

    for fragment, unit in zip(
        ("лобненск", "петрозаводск", "планетн"),
        UNITS,
    ):
        if fragment in key:
            return unit, unit

    match = re.search(
        r"(?:женская консультация|жк)\s*(\d+)\b",
        key,
    )

    if match:
        n = int(match[1])

        if n in (3, 5, 6, 7):
            return PLANET, f"ЖК №{n}"

        if n == 10:
            return UNITS[3], UNITS[3]

    return None


@dataclass
class Bucket:
    quantity: Decimal = ZERO
    money: Decimal = ZERO
    doctors: set = field(default_factory=set)
    codes: set = field(default_factory=set)
    positions: set = field(default_factory=set)

    def add(self, doctor, code, name, tariff, quantity):
        self.quantity += quantity

        self.money += (
            tariff * quantity
        ).quantize(CENT, rounding=ROUND_HALF_UP)

        self.doctors.add(doctor)
        self.codes.add(code)

        self.positions.add(
            (doctor, code, normalize_key(name), tariff)
        )


@dataclass
class PeriodData:
    units: dict = field(
        default_factory=lambda: defaultdict(Bucket)
    )
    services: dict = field(
        default_factory=lambda: defaultdict(Bucket)
    )
    doctors: dict = field(
        default_factory=lambda: defaultdict(Bucket)
    )
    sources: dict = field(
        default_factory=lambda: defaultdict(Bucket)
    )
    total: Bucket = field(default_factory=Bucket)
    code_totals: dict = field(
        default_factory=lambda: defaultdict(Bucket)
    )
    names: dict = field(default_factory=dict)
    doctor_names: dict = field(default_factory=dict)
    dates: Counter = field(default_factory=Counter)
    excluded: Counter = field(default_factory=Counter)
    empty_services: int = 0
    outside: int = 0
    restored_tariffs: Counter = field(default_factory=Counter)
    day_hospital: PeriodData | None = None
    hospital_source: str = ""
    hospital_sheet: str = ""
    hospital_stats: object | None = None

def collect_tariffs(sheet, header, epoch, start, end):
    """Собирает заполненные тарифы нужных ЖК за выбранный период."""
    c = header.columns
    tariffs = defaultdict(set)

    for row in sheet.iter_rows(min_row=header.row_number + 1):
        try:
            if not identify_department(
                clean_text(row[c.department].value)
            ):
                continue

            dt = parse_date(row[c.service_date].value, epoch=epoch)
            if not start <= dt <= end:
                continue

            code = code_key(
                service_code_from_cell(row[c.service_code])
            )
            raw_tariff = row[c.tariff].value

            if not code or not clean_text(raw_tariff):
                continue

            tariff = number(raw_tariff)
            if tariff >= 0:
                tariffs[code].add(tariff)

        except (ValueError, TypeError, IndexError):
            continue

    return tariffs


def resolve_tariff(raw_tariff, code, tariffs):
    """Заполняет только пустой тариф при однозначном совпадении кода."""
    if clean_text(raw_tariff):
        return number(raw_tariff)

    options = tariffs.get(code, set())

    if not options:
        raise ValueError(
            f"пустой тариф для услуги {code}: "
            "в выбранном периоде нет заполненного тарифа этого кода"
        )

    if len(options) > 1:
        variants = ", ".join(str(value) for value in sorted(options))
        raise ValueError(
            f"пустой тариф для услуги {code}: "
            f"найдены разные тарифы ({variants}); "
            "однозначная подстановка невозможна"
        )

    return next(iter(options))


def read_period(path, start, end, day_hospital_path=None):
    path = Path(path)

    if path.suffix.lower() not in (".xls", ".xlsx", ".xlsm"):
        raise ReportError(
            "Поддерживаются .xls, .xlsx и .xlsm."
        )

    if path.suffix.lower() == ".xls":
        book = XlsWorkbookAdapter(path)
    else:
        book = load_workbook(
            path,
            read_only=True,
            data_only=True,
            keep_links=False,
        )

    data = PeriodData()

    try:
        resolver = HeaderResolver(AppConfig())
        candidates = []

        for sheet in book.worksheets:
            try:
                candidates.append(
                    (sheet, resolver.find(sheet))
                )
            except ReportError:
                continue

        if len(candidates) != 1:
            raise ReportError(
                "В исходнике должен быть ровно один лист "
                "с таблицей услуг."
            )

        sheet, header = candidates[0]
        c = header.columns

        tariffs = collect_tariffs(sheet, header, book.epoch, start, end)

        for row_no, row in enumerate(
            sheet.iter_rows(
                min_row=header.row_number + 1
            ),
            header.row_number + 1,
        ):
            values = [cell.value for cell in row]

            if all(
                v is None or v == ""
                for v in values
            ):
                continue

            try:
                dep = clean_text(values[c.department])

                if not dep:
                    raise ValueError(
                        "не заполнено подразделение"
                    )

                match = identify_department(dep)

                if not match:
                    key = normalize_key(dep)

                    if (
                        "женск" in key
                        or re.search(r"\b(?:жк|цжз)\b", key)
                    ):
                        dt = parse_date(
                            values[c.service_date],
                            epoch=book.epoch,
                        )

                        if start <= dt <= end:
                            data.excluded[dep] += 1

                    continue

                dt = parse_date(
                    values[c.service_date],
                    epoch=book.epoch,
                )

                if not start <= dt <= end:
                    data.outside += 1
                    continue

                raw = [
                    values[i]
                    for i in (
                        c.service_code,
                        c.service_name,
                        c.quantity,
                        c.tariff,
                    )
                ]

                if all(
                    v is None or v == ""
                    for v in raw
                ):
                    data.empty_services += 1
                    continue

                doctor_name = safe_excel_text(
                    values[c.doctor]
                )
                name = safe_excel_text(
                    values[c.service_name]
                )
                code = code_key(
                    service_code_from_cell(
                        row[c.service_code]
                    )
                )

                if not doctor_name or not name or not code:
                    raise ValueError(
                        "не заполнены врач, код "
                        "или название услуги"
                    )

                quantity = number(values[c.quantity])
                raw_tariff = values[c.tariff]
                tariff = resolve_tariff(raw_tariff, code, tariffs)

                if not clean_text(raw_tariff):
                    data.restored_tariffs[(code, tariff)] += 1

                if (
                    quantity < 0
                    or quantity != quantity.to_integral_value()
                    or tariff < 0
                ):
                    raise ValueError(
                        "ожидались целое количество >= 0 "
                        "и тариф >= 0"
                    )

                doctor = normalize_key(doctor_name)
                unit, source = match

                buckets = (
                    data.units[unit],
                    data.services[(unit, code)],
                    data.doctors[(unit, doctor)],
                    data.sources[(unit, source)],
                    data.total,
                    data.code_totals[code],
                )

                for bucket in buckets:
                    bucket.add(
                        doctor,
                        code,
                        name,
                        tariff,
                        quantity,
                    )

                data.names[code] = name
                data.doctor_names[doctor] = doctor_name
                data.dates[dt] += 1

            except (
                ValueError,
                TypeError,
                IndexError,
            ) as exc:
                raise ReportError(
                    f"{path.name}, лист {sheet.title}, "
                    f"строка {row_no}: {exc}"
                ) from exc

        if day_hospital_path is not None:
            attach_day_hospital(data, read_day_hospital(day_hospital_path, start, end))

        if not data.dates:
            raise ReportError(
                f"{path.name}: нет услуг нужных "
                "подразделений за выбранный период."
            )

        return data

    finally:
        book.close()


def read_day_hospital(path, start, end):
    """Использует общие с zhk_report правила дат, суммы и повторов."""
    try:
        from zhk_report import DayHospitalReader
    except ImportError as exc:
        raise ReportError("Для стационара замените также zhk_report.py из архива обновления.") from exc
    path = Path(path)
    if not path.is_file() or path.suffix.lower() not in (".xls", ".xlsx", ".xlsm"):
        raise ReportError(f"Выберите Excel-файл дневного стационара: {path}")
    departments, stats, sheet = DayHospitalReader(AppConfig()).aggregate(
        path, start, end, allow_erzl_fallback=True,
    )
    data = PeriodData(hospital_source=path.name, hospital_sheet=sheet, hospital_stats=stats)
    for department in departments:
        unit, _ = identify_department(department.name)
        source = "Дневной стационар №1-1" if unit == UNITS[0] else "Дневной стационар №2-1"
        for record in department.sorted_records():
            doctor_name = safe_excel_text(record.doctor)
            doctor = normalize_key(doctor_name)
            code = code_key(record.service_code)
            name = safe_excel_text(record.service_name)
            for bucket in (
                data.units[unit], data.services[(unit, code)], data.doctors[(unit, doctor)],
                data.sources[(unit, source)], data.total, data.code_totals[code],
            ):
                bucket.add(doctor, code, name, record.tariff, record.quantity)
            data.names[code] = name
            data.doctor_names[doctor] = doctor_name
        data.dates.update({day: int(qty) for day, qty in department.daily_quantity.items()})
    return data


def merge_bucket(target, addition):
    target.quantity += addition.quantity
    target.money += addition.money
    target.doctors.update(addition.doctors)
    target.codes.update(addition.codes)
    target.positions.update(addition.positions)


def attach_day_hospital(data, hospital):
    if data.day_hospital is not None:
        raise ReportError("Дневной стационар уже включён в этот период.")
    for field_name in ("units", "services", "doctors", "sources", "code_totals"):
        target = getattr(data, field_name)
        for key, bucket in getattr(hospital, field_name).items():
            merge_bucket(target[key], bucket)
    merge_bucket(data.total, hospital.total)
    data.dates.update(hospital.dates)
    for code, name in hospital.names.items():
        data.names.setdefault(code, name)
    data.doctor_names.update(hospital.doctor_names)
    data.day_hospital = hospital


def validate_hospital_files(paths, main_paths):
    """Оба периода должны иметь одинаковый состав источников."""
    if bool(paths[0]) != bool(paths[1]):
        raise ReportError("Выберите оба файла стационара или очистите оба поля стационара.")
    protected = {Path(p).resolve() for p in main_paths}
    for path in paths:
        if path is None:
            continue
        if not path.is_file() or path.suffix.lower() not in (".xls", ".xlsx", ".xlsm"):
            raise ReportError(f"Не найден Excel-файл дневного стационара: {path}")
        if path.resolve() in protected:
            raise ReportError("Для стационара нужна отдельная выгрузка, отличная от услуг ЖК и шаблона.")


def validate_periods(a, b, c, d):
    for start, end in ((a, b), (c, d)):
        if start > end:
            raise ReportError(
                "Дата начала периода не может быть позже даты окончания."
            )

        if (start.year, start.month) != (end.year, end.month):
            raise ReportError(
                "Каждый период должен находиться внутри одного месяца."
            )

    if c.year * 12 + c.month - (a.year * 12 + a.month) != 1:
        raise ReportError(
            "Нужны предыдущий и следующий за ним месяц."
        )


def pair(mapping0, mapping1, key):
    return (
        mapping0.get(key, Bucket()),
        mapping1.get(key, Bucket()),
    )


def metrics(prefix, a, b, row):
    """Количество, изменение, суммы и средние."""
    n = len(prefix)

    def col(offset):
        return f"{get_column_letter(n + offset)}{row}"

    (
        q0, q1, delta, pct,
        m0, m1, dm,
        av0, av1, dav,
    ) = map(col, range(1, 11))

    return prefix + [
        a.quantity,
        b.quantity,
        f"={q1}-{q0}",
        f'=IF({q0}=0,"",{delta}/{q0})',
        a.money,
        b.money,
        f"={m1}-{m0}",
        f'=IF({q0}=0,"",{m0}/{q0})',
        f'=IF({q1}=0,"",{m1}/{q1})',
        (
            f'=IF(OR({av0}="",{av1}=""),'
            f'"",{av1}-{av0})'
        ),
    ]


class TemplateWriter:
    def __init__(self, path):
        self.book = load_workbook(path)
        self.styles = {}
        self.headers = {}
        self.groups = {}

        expected = {
            "Коротко": (
                "ShortSummary14",
                "OverallGrowth14",
                "OverallDecline14",
            ),
            "Подразделения": (
                "UnitsComparison14",
            ),
            "Главные изменения": (
                "TopServiceChanges14",
            ),
            "Все услуги": (
                "AllServices14",
            ),
            "По врачам": (
                "DoctorComparison14",
            ),
            "Планетная": (
                "PlanetaSummary14",
                "PlanetaSources14",
                "PlanetaGroups14",
            ),
        }

        if self.book.sheetnames != list(expected):
            raise ReportError(
                "Нужен исходный шаблон "
                "с шестью листами из примера."
            )

        for row in self.book["Все услуги"].iter_rows(
            min_row=6,
            values_only=True,
        ):
            if row[2] is not None:
                code = code_key(row[2])
                group = row[1]

                if (
                    code in self.groups
                    and self.groups[code] != group
                ):
                    raise ReportError(
                        "В шаблоне разные группы "
                        f"для кода {code}."
                    )

                self.groups[code] = group

        for sheet_name, names in expected.items():
            ws = self.book[sheet_name]

            if set(ws.tables) != set(names):
                raise ReportError(
                    "Изменилась структура таблиц "
                    f"шаблона: {sheet_name}."
                )

            for name in names:
                table = ws.tables[name]
                c1, r1, c2, _ = range_boundaries(
                    table.ref
                )

                self.headers[name] = [
                    ws.cell(r1, c).value
                    for c in range(c1, c2 + 1)
                ]

                self.styles[name] = (
                    [
                        copy(ws.cell(r1, c)._style)
                        for c in range(c1, c2 + 1)
                    ],
                    [
                        copy(ws.cell(r1 + 1, c)._style)
                        for c in range(c1, c2 + 1)
                    ],
                    copy(table.tableStyleInfo),
                )

            ws.tables.clear()
            ws.conditional_formatting = (
                ConditionalFormattingList()
            )

            for row in ws:
                for cell in row:
                    if not isinstance(cell, MergedCell):
                        cell.value = None
                        cell.hyperlink = None
                        cell.comment = None

            ws.sheet_view.showGridLines = False
            ws.sheet_view.zoomScale = 85
            ws.print_options.horizontalCentered = False
            ws.page_setup.orientation = "landscape"
            ws.page_setup.paperSize = ws.PAPERSIZE_A3
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
            ws.sheet_properties.pageSetUpPr.fitToPage = True

        # Блок источников Планетной может менять размер.
        for merged in list(
            self.book["Планетная"].merged_cells.ranges
        ):
            if merged.min_row >= 11:
                self.book["Планетная"].unmerge_cells(
                    str(merged)
                )

        self.ends = defaultdict(lambda: 1)

    def text(self, sheet, address, value):
        self.book[sheet][address] = value

    def table(
        self,
        sheet,
        name,
        row,
        col,
        rows,
        labels,
    ):
        ws = self.book[sheet]
        hs, ds, table_style = self.styles[name]

        rows = rows or [
            ["Нет данных"] + [None] * (len(labels) - 1)
        ]

        for offset, values in enumerate(
            [labels] + rows
        ):
            rr = row + offset
            ws.row_dimensions[rr].height = (
                64 if offset == 0 else 46
            )

            for j, value in enumerate(values):
                cell = ws.cell(rr, col + j)
                cell._style = copy(
                    (hs if offset == 0 else ds)[j]
                )

                is_text = (
                    isinstance(value, str)
                    and not value.startswith("=")
                )

                cell.alignment = Alignment(
                    vertical="center",
                    wrap_text=True,
                    horizontal="left" if is_text else "right",
                )

                if isinstance(value, Decimal):
                    value = (
                        int(value)
                        if value == value.to_integral_value()
                        else float(value)
                    )

                cell.value = value

                if (
                    isinstance(value, str)
                    and not value.startswith("=")
                ):
                    cell.data_type = "s"

        end = row + len(rows)

        ref = (
            f"{get_column_letter(col)}{row}:"
            f"{get_column_letter(col + len(labels) - 1)}"
            f"{end}"
        )

        ws.add_table(
            Table(
                displayName=name,
                ref=ref,
                tableStyleInfo=table_style,
            )
        )

        self.ends[sheet] = max(
            self.ends[sheet],
            end,
        )

        return end

    def colors(self, sheet, columns, start, end):
        if end < start:
            return

        for col in columns:
            for operator, fill, color in (
                ("greaterThan", "E2F0D9", "216E39"),
                ("lessThan", "FCE4D6", "A12622"),
            ):
                self.book[
                    sheet
                ].conditional_formatting.add(
                    f"{col}{start}:{col}{end}",
                    CellIsRule(
                        operator=operator,
                        formula=["0"],
                        fill=PatternFill(
                            "solid",
                            fgColor=fill,
                        ),
                        font=Font(color=color),
                    ),
                )

    def finish(self):
        for ws in self.book:
            last = max(
                self.ends[ws.title],
                24 if ws.title == "Коротко" else 1,
            )

            last_col = max(
                range_boundaries(t.ref)[2]
                for t in ws.tables.values()
            )

            ws.print_area = (
                f"A1:{get_column_letter(last_col)}{last}"
            )

            ws.sheet_view.topLeftCell = "A1"
            pane = ws.sheet_view.pane

            ws.sheet_view.selection = [
                Selection(
                    pane=pane.activePane if pane else None,
                    activeCell=ws.freeze_panes or "A1",
                    sqref=ws.freeze_panes or "A1",
                )
            ]

        self.book.calculation = CalcProperties(
            calcId=0,
            fullCalcOnLoad=True,
            forceFullCalc=True,
            calcMode="auto",
        )

        return self.book


def period_header(value, labels):
    # Одновременная замена: «Сентябрь» из нового первого заголовка
    # не должен повторно превратиться в «Октябрь».
    return re.sub(
        r"август|сентябрь",
        lambda match: labels[0] if match.group().casefold() == "август" else labels[1],
        str(value), flags=re.IGNORECASE,
    )


def write_day_hospital_sheet(writer, old, new, labels, period):
    """Отдельная расшифровка уже включённых в ЖК услуг; не новый общий итог."""
    old, new = old.day_hospital, new.day_hospital
    name = "Дневной стационар"
    ws = writer.book.create_sheet(name, writer.book.sheetnames.index("Планетная") + 1)
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
    ws.sheet_properties.tabColor = "1F4E78"
    ws.freeze_panes = "B6"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:5"
    widths = (34, 24, 18, 28, 18, 20, 20, 18, 22, 22, 22, 22, 22, 22)
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width

    def title(row, text, size=12):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=14)
        cell = ws.cell(row, 1, text)
        cell.font = Font(name="Calibri", size=size, bold=True, color="1F4E78")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 28 if size < 16 else 34

    title(1, "Дневной стационар: сравнение периодов", 18)
    title(2, period, 11)
    title(3, "№1-1 включён в Лобненскую; №2-1 — в Петрозаводскую. Ниже — расшифровка этих сумм.", 11)

    def labels_for(base, prefix):
        return prefix + [
            period_header(v, labels)
            for v in writer.headers[base][len(prefix):]
        ]

    def write_table(table_name, base, header_row, rows, headers):
        writer.styles[table_name] = writer.styles[base]
        end = writer.table(name, table_name, header_row, 1, rows, headers)
        average0, average1, delta = {
            "PlanetaGroups14": ("I", "J", "K"),
            "AllServices14": ("L", "M", "N"),
            "DoctorComparison14": ("J", "K", "L"),
        }[base]
        for row in range(header_row + 1, end + 1):
            ws[f"{delta}{row}"] = (
                f'=IF(OR({average0}{row}="",{average1}{row}=""),"",'
                f'ROUND({average1}{row}-{average0}{row},2))'
            )
        return end

    summary = []
    for row, unit, label in (
        (6, UNITS[0], "ДС №1-1 (Лобненская)"),
        (7, UNITS[1], "ДС №2-1 (Петрозаводская)"),
    ):
        summary.append(metrics([label], *pair(old.units, new.units, unit), row))
    total = metrics(["ИТОГО по стационару"], old.total, new.total, 8)
    for index, col in ((1, "B"), (2, "C"), (5, "F"), (6, "G")):
        total[index] = f"=SUM({col}6:{col}7)"
    summary.append(total)
    write_table("DayHospitalSummary", "PlanetaGroups14", 5, summary,
                labels_for("PlanetaGroups14", ["Стационар / ЖК"]))
    writer.colors(name, ("D", "H", "K"), 6, 8)
    for cell in ws[8][:11]:
        cell.font = Font(name="Calibri", size=11, bold=True, color="1F4E78")

    for row, label, data in ((10, labels[0], old), (11, labels[1], new)):
        stats = data.hospital_stats
        detail = (
            f"{label}: {data.hospital_source}; принято {stats.accepted_rows}, "
            f"повторов исключено {stats.duplicate_rows}; проверка по ЕРЗЛ при пустом полисе: {stats.erzl_rows}."
        )
        if stats.erzl_examples:
            detail += " Строки: " + ", ".join(map(str, stats.erzl_examples)) + "."
        if stats.without_discharge_date:
            detail += f" Без даты выбытия: {stats.without_discharge_date} (не включены)."
        title(row, detail, 10)
        ws.cell(row, 1).font = Font(name="Calibri", size=10, color="595959")

    title(13, "По кодам МЭС")
    names = old.names | new.names
    keys = sorted(set(old.services) | set(new.services), key=lambda k: (UNITS.index(k[0]), k[1]))
    rows = []
    for row, (unit, code) in enumerate(keys, 15):
        station = "№1-1" if unit == UNITS[0] else "№2-1"
        rows.append(metrics([unit, station, code, names[code]],
                            *pair(old.services, new.services, (unit, code)), row))
    end = write_table("DayHospitalMES", "AllServices14", 14, rows,
                      labels_for("AllServices14", ["ЖК", "Стационар", "Код МЭС", "МЭС. Название"]))
    writer.colors(name, ("G", "K", "N"), 15, end)
    for row in range(15, end + 1):
        ws.cell(row, 3).number_format = "@"
        ws.row_dimensions[row].height = max(46, 16 * ((len(str(ws.cell(row, 4).value or "")) + 24) // 25))

    doctor_header = end + 4
    title(doctor_header - 1, "По врачам дневного стационара")
    doctor_names = old.doctor_names | new.doctor_names
    keys = sorted(set(old.doctors) | set(new.doctors), key=lambda k: (UNITS.index(k[0]), k[1]))
    rows = []
    for row, key in enumerate(keys, doctor_header + 1):
        unit, doctor = key
        x, y = pair(old.doctors, new.doctors, key)
        rows.append(metrics([unit, doctor_names[doctor]], x, y, row) + [len(x.codes), len(y.codes)])
    end = write_table("DayHospitalDoctors", "DoctorComparison14", doctor_header, rows,
                      labels_for("DoctorComparison14", ["ЖК", "Врач"]))
    writer.colors(name, ("E", "I", "L"), doctor_header + 1, end)


def build_book(template, old, new, dates):
    if (old.day_hospital is None) != (new.day_hospital is None):
        raise ReportError("Для сравнения со стационаром загрузите его выгрузки для обоих периодов.")
    writer = TemplateWriter(template)
    if old.day_hospital is not None:
        for code in set(old.day_hospital.names) | set(new.day_hospital.names):
            writer.groups.setdefault(code, "Дневной стационар")
    a, b, c, d = dates

    labels = (
        f"{MONTHS[a.month]} {a.year}",
        f"{MONTHS[c.month]} {c.year}",
    )

    period = (
        f"{a:%d.%m.%Y}–{b:%d.%m.%Y} и "
        f"{c:%d.%m.%Y}–{d:%d.%m.%Y}"
    )

    names = old.names | new.names
    doctors = old.doctor_names | new.doctor_names

    def headers(name):
        return [
            period_header(v, labels)
            for v in writer.headers[name]
        ]

    def group(code):
        return safe_excel_text(
            writer.groups.get(
                code,
                "Не классифицировано",
            )
        )

    def table(sheet, name, row, col, rows):
        return writer.table(
            sheet,
            name,
            row,
            col,
            rows,
            headers(name),
        )

    def ranks(keys, get_pair, sign, limit):
        matching = [
            k
            for k in keys
            if (
                get_pair(k)[1].quantity
                - get_pair(k)[0].quantity
            ) * sign > 0
        ]

        return sorted(
            matching,
            key=lambda k: (
                -sign * (
                    get_pair(k)[1].quantity
                    - get_pair(k)[0].quantity
                ),
                str(k),
            ),
        )[:limit]

    service_keys = (
        set(old.services) | set(new.services)
    )

    ordered = sorted(
        service_keys,
        key=lambda k: (
            UNITS.index(k[0]),
            -(
                new.services.get(k, Bucket()).quantity
                - old.services.get(k, Bucket()).quantity
            ),
            k[1],
        ),
    )

    all_rows = []
    service_row = {}

    for r, (unit, code) in enumerate(ordered, 6):
        service_row[(unit, code)] = r

        x, y = pair(
            old.services,
            new.services,
            (unit, code),
        )

        all_rows.append(
            metrics(
                [unit, group(code), code, names[code]],
                x,
                y,
                r,
            )
        )

    writer.text(
        "Все услуги",
        "A1",
        "Полная расшифровка по кодам услуг",
    )
    writer.text(
        "Все услуги",
        "A2",
        period
        + ". Разница = текущий период минус предыдущий.",
    )

    end = table(
        "Все услуги",
        "AllServices14",
        5,
        1,
        all_rows,
    )

    writer.book["Все услуги"].freeze_panes = "E6"
    writer.colors(
        "Все услуги",
        ("G", "K", "N"),
        6,
        end,
    )

    unit_rows = []

    for r, unit in enumerate((*UNITS, "ИТОГО"), 6):
        if unit == "ИТОГО":
            x, y = old.total, new.total
        else:
            x, y = pair(
                old.units,
                new.units,
                unit,
            )

        row = metrics([unit], x, y, r) + [
            len(x.doctors),
            len(y.doctors),
            len(x.codes),
            len(y.codes),
            len(y.codes - x.codes),
            len(x.codes - y.codes),
        ]

        unit_rows.append(row)

    writer.text(
        "Подразделения",
        "A1",
        "Сравнение подразделений за выбранные периоды",
    )
    writer.text(
        "Подразделения",
        "A2",
        period + ". Средняя стоимость = сумма / количество.",
    )

    table(
        "Подразделения",
        "UnitsComparison14",
        5,
        1,
        unit_rows,
    )

    writer.colors(
        "Подразделения",
        ("D", "H", "K"),
        6,
        10,
    )

    changes = []

    for unit in UNITS:
        keys = [
            k
            for k in service_keys
            if k[0] == unit
        ]

        def get_pair(k):
            return pair(
                old.services,
                new.services,
                k,
            )

        for sign, direction in (
            (1, "Рост"),
            (-1, "Снижение"),
        ):
            for key in ranks(
                keys,
                get_pair,
                sign,
                7,
            ):
                rr = service_row[key]

                changes.append(
                    [
                        unit,
                        direction,
                        key[1],
                        names[key[1]],
                    ]
                    + [
                        f"='Все услуги'!{col}{rr}"
                        for col in (
                            "E", "F", "G",
                            "I", "J", "K",
                            "L", "M",
                        )
                    ]
                )

    writer.text(
        "Главные изменения",
        "A1",
        "Какие услуги изменились сильнее всего",
    )
    writer.text(
        "Главные изменения",
        "A2",
        "До 7 ростов и 7 снижений по количеству "
        "в каждом подразделении.",
    )

    end = table(
        "Главные изменения",
        "TopServiceChanges14",
        5,
        1,
        changes,
    )

    writer.colors(
        "Главные изменения",
        ("G", "J"),
        6,
        end,
    )
    writer.book[
        "Главные изменения"
    ].freeze_panes = "E6"

    doctor_keys = (
        set(old.doctors) | set(new.doctors)
    )

    doctor_keys = sorted(
        doctor_keys,
        key=lambda k: (
            UNITS.index(k[0]),
            -abs(
                new.doctors.get(k, Bucket()).quantity
                - old.doctors.get(k, Bucket()).quantity
            ),
            k[1],
        ),
    )

    rows = []

    for r, key in enumerate(doctor_keys, 5):
        x, y = pair(
            old.doctors,
            new.doctors,
            key,
        )

        rows.append(
            metrics(
                [key[0], doctors[key[1]]],
                x,
                y,
                r,
            )
            + [
                len(x.codes),
                len(y.codes),
            ]
        )

    writer.text(
        "По врачам",
        "A1",
        "Сравнение по врачам",
    )

    end = table(
        "По врачам",
        "DoctorComparison14",
        4,
        1,
        rows,
    )

    writer.colors(
        "По врачам",
        ("E", "I", "L"),
        5,
        end,
    )
    writer.book["По врачам"].freeze_panes = "C5"

    writer.text(
        "Коротко",
        "A1",
        "Сравнение двух периодов",
    )
    writer.text(
        "Коротко",
        "A2",
        period + ". Главный показатель — количество услуг.",
    )
    writer.text(
        "Коротко",
        "A4",
        "Планетная: ЖК №3, №5, №6, №7 "
        "и ЦЖЗ на Планетной объединены.",
    )

    rows = [
        [unit]
        + [
            f"='Подразделения'!"
            f"{get_column_letter(col)}{r}"
            for col in range(2, 14)
        ]
        for r, unit in enumerate(
            (*UNITS, "ИТОГО"),
            6,
        )
    ]

    table(
        "Коротко",
        "ShortSummary14",
        6,
        1,
        rows,
    )

    writer.colors(
        "Коротко",
        ("D", "H", "K"),
        7,
        11,
    )

    keys = (
        set(old.code_totals) | set(new.code_totals)
    )

    def get_pair(k):
        return pair(
            old.code_totals,
            new.code_totals,
            k,
        )

    for sign, name, col, cell, title in (
        (
            1,
            "OverallGrowth14",
            1,
            "A14",
            "Крупнейший рост по кодам",
        ),
        (
            -1,
            "OverallDecline14",
            8,
            "H14",
            "Крупнейшее снижение по кодам",
        ),
    ):
        rows = []

        for code in ranks(
            keys,
            get_pair,
            sign,
            6,
        ):
            x, y = get_pair(code)

            rows.append([
                code,
                names[code],
                x.quantity,
                y.quantity,
                y.quantity - x.quantity,
                y.money - x.money,
            ])

        writer.text(
            "Коротко",
            cell,
            title,
        )

        table(
            "Коротко",
            name,
            15,
            col,
            rows,
        )

    writer.text(
        "Коротко",
        "A24",
        "Зелёный — рост; красный — снижение.",
    )
    writer.book["Коротко"].freeze_panes = "B7"

    writer.text(
        "Планетная",
        "A1",
        "ЖК №3, №5, №6, №7 и ЦЖЗ на Планетной",
    )
    writer.text(
        "Планетная",
        "A2",
        period,
    )

    summary = []

    for r, title, source_col in (
        (6, "Количество услуг", "B"),
        (7, "Сумма, руб.", "F"),
        (8, "Средняя стоимость 1 услуги, руб.", "I"),
    ):
        second_col = get_column_letter(
            range_boundaries(
                f"{source_col}1:{source_col}1"
            )[0] + 1
        )

        summary.append([
            title,
            f"='Подразделения'!{source_col}8",
            f"='Подразделения'!{second_col}8",
            f"=C{r}-B{r}",
            f'=IF(B{r}=0,"",D{r}/B{r})',
        ])

    table(
        "Планетная",
        "PlanetaSummary14",
        5,
        1,
        summary,
    )

    ws = writer.book["Планетная"]
    ws["B5"], ws["C5"] = labels

    writer.text(
        "Планетная",
        "A11",
        "Какие источники вошли в сравнение",
    )

    sources = []

    for label, data in zip(labels, (old, new)):
        for (unit, source), bucket in sorted(
            data.sources.items()
        ):
            if unit == PLANET:
                r = 13 + len(sources)

                sources.append([
                    label,
                    "ЦЖЗ" if source == PLANET else "Отдельные ЖК",
                    source,
                    bucket.quantity,
                    bucket.money,
                    f'=IF(D{r}=0,"",E{r}/D{r})',
                    len(bucket.doctors),
                    len(bucket.codes),
                    len(bucket.positions),
                ])

    end = table(
        "Планетная",
        "PlanetaSources14",
        12,
        1,
        sources,
    )

    ws["I12"] = "Позиций врач–код–название–тариф"

    group_header = max(21, end + 4)

    writer.text(
        "Планетная",
        f"A{group_header - 1}",
        "Какие группы изменили объём Планетной",
    )

    group_rows = []

    groups = sorted(
        {
            group(code)
            for unit, code in service_keys
            if unit == PLANET
        },
        key=lambda g: (
            g == "Другие услуги",
            g,
        ),
    )

    for r, g in enumerate(
        groups,
        group_header + 1,
    ):
        aggregates = []

        for data in (old, new):
            entries = [
                v
                for (unit, code), v in data.services.items()
                if unit == PLANET and group(code) == g
            ]

            aggregates.append(
                Bucket(
                    quantity=sum(
                        (v.quantity for v in entries),
                        ZERO,
                    ),
                    money=sum(
                        (v.money for v in entries),
                        ZERO,
                    ),
                )
            )

        group_rows.append(
            metrics(
                [g],
                *aggregates,
                r,
            )
        )

    end = table(
        "Планетная",
        "PlanetaGroups14",
        group_header,
        1,
        group_rows,
    )

    writer.colors(
        "Планетная",
        ("D",),
        6,
        8,
    )
    writer.colors(
        "Планетная",
        ("D", "H", "K"),
        group_header + 1,
        end,
    )

    ws.freeze_panes = "B6"

    if old.day_hospital is not None:
        write_day_hospital_sheet(writer, old, new, labels, period)

    unknown = sorted(
        set(names) - set(writer.groups)
    )

    notes = [
        (
            "Записи предыдущего периода: "
            f"{min(old.dates):%d.%m.%Y}–"
            f"{max(old.dates):%d.%m.%Y}; "
            f"дней с услугами {len(old.dates)} из {(b - a).days + 1}."
        ),
        (
            "Записи текущего периода: "
            f"{min(new.dates):%d.%m.%Y}–"
            f"{max(new.dates):%d.%m.%Y}; "
            f"дней с услугами {len(new.dates)} из {(d - c).days + 1}."
        ),
        (
            "Строк без данных услуги: "
            f"{old.empty_services} / {new.empty_services}. "
            "Врачей: уникальные нормализованные ФИО; "
            "совпадающие ФИО не различаются."
        ),
    ]

    notes.append(
        "Автоматически заполнено пустых тарифов: "
        f"{sum(old.restored_tariffs.values())} в предыдущем периоде; "
        f"{sum(new.restored_tariffs.values())} в текущем."
    )
    if old.day_hospital is not None:
        notes.append(
            "Дневной стационар включён в общие итоги: №1-1 — Лобненская, №2-1 — Петрозаводская. "
            "Отдельный лист «Дневной стационар» содержит расшифровку уже учтённых услуг."
        )

    for label, data in zip(labels, (old, new)):
        if data.excluded:
            notes.append(
                label
                + "; вне периметра: "
                + "; ".join(
                    f"{name}: {count} строк"
                    for name, count in data.excluded.items()
                )
            )

    if unknown:
        notes.append(
            "Не классифицированы коды: "
            + ", ".join(unknown)
        )

    ws = writer.book["Коротко"]

    for r, note in enumerate(notes, 26):
        ws.merge_cells(
            start_row=r,
            start_column=1,
            end_row=r,
            end_column=13,
        )
        ws.cell(r, 1, note)
        ws.cell(r, 1).font = Font(
            name="Calibri",
            size=10,
            color="595959",
        )
        ws.cell(r, 1).alignment = Alignment(
            wrap_text=True,
            vertical="center",
        )
        ws.row_dimensions[r].height = 42
        writer.ends["Коротко"] = r

    return writer.finish(), unknown


def describe(data, label, unknown=()):
    days = sorted(data.dates)

    parts = [
        (
            f"{label}: {data.total.quantity} услуг; "
            f"{data.total.money:,.2f} руб.; "
            f"{len(data.total.doctors)} уникальных ФИО."
        ),
        (
            "Даты записей: "
            f"{days[0]:%d.%m.%Y}–"
            f"{days[-1]:%d.%m.%Y}; "
            f"дней с услугами: {len(days)}."
        ),
        (
            "Строк без данных услуги: "
            f"{data.empty_services}. "
            "Строк нужных подразделений вне периода: "
            f"{data.outside}."
        ),
    ]

    if data.excluded:
        parts.append(
            "Вне периметра образца: "
            + "; ".join(
                f"{name}: {n} строк"
                for name, n in data.excluded.items()
            )
        )

    if unknown:
        parts.append(
            "Нет группы в шаблоне: "
            + ", ".join(unknown)
        )

        if data.restored_tariffs:
            parts.append(
                "Автоматически заполнены пустые тарифы: "
                + "; ".join(
                    f"код {code}: {tariff} руб., строк: {count}"
                    for (code, tariff), count
                    in sorted(data.restored_tariffs.items())
                )
            )

    if data.day_hospital is not None:
        hospital = data.day_hospital
        stats = hospital.hospital_stats
        parts.append(
            f"В том числе дневной стационар: {hospital.total.quantity} услуг; "
            f"{hospital.total.money:,.2f} руб. ({hospital.hospital_source}, лист «{hospital.hospital_sheet}»). "
            f"Повторов исключено: {stats.duplicate_rows}; строк с проверкой по ЕРЗЛ: {stats.erzl_rows}."
        )
        if stats.erzl_examples:
            parts.append("Пустой полис, использован ЕРЗЛ: строки " + ", ".join(map(str, stats.erzl_examples)) + ".")
        parts.extend(stats.duplicate_examples)
        if stats.without_discharge_date:
            parts.append(f"Без даты выбытия, не включены в период: {stats.without_discharge_date}.")

    return "\n".join(parts)


def _comparison_span(chosen, edge="start"):
    """Ровно 14 дней включительно, внутри одного месяца."""
    if edge not in ("start", "end"):
        raise ValueError("Неизвестная граница периода.")

    start = (
        chosen
        if edge == "start"
        else chosen - timedelta(days=13)
    )
    end = start + timedelta(days=13)

    if (start.year, start.month) != (end.year, end.month):
        raise ValueError(
            "Эти 14 дней выходят за границу месяца. "
            "Выберите начало раньше или окончание позже."
        )

    return start, end


def launch_compare(parent=None):
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk
    from zhk_gui import CalendarPopup

    window = tk.Toplevel(parent) if parent else tk.Tk()
    window.title("Сравнение услуг за выбранные периоды")
    window.withdraw()

    width = min(1020, window.winfo_screenwidth() - 60)
    height = min(780, window.winfo_screenheight() - 100)

    window.geometry(f"{width}x{height}")
    window.minsize(min(860, width), min(600, height))

    style = ttk.Style(window)

    if parent is None and "clam" in style.theme_names():
        style.theme_use("clam")

    style.configure(
        "Compare.TFrame",
        background="#F4F7FB",
    )
    style.configure(
        "CompareCard.TFrame",
        background="white",
    )
    style.configure(
        "Compare.TLabel",
        background="white",
        foreground="#243B53",
        font=("Segoe UI", 10),
    )
    style.configure(
        "CompareTitle.TLabel",
        background="#F4F7FB",
        foreground="#173B5E",
        font=("Segoe UI", 18, "bold"),
    )
    style.configure(
        "CompareHeading.TLabel",
        background="white",
        foreground="#173B5E",
        font=("Segoe UI", 11, "bold"),
    )
    style.configure(
        "CompareHint.TLabel",
        background="#F4F7FB",
        foreground="#526779",
        font=("Segoe UI", 9),
    )
    style.configure(
        "Compare.TEntry",
        padding=6,
    )
    style.configure(
        "Compare.TButton",
        padding=(10, 7),
        font=("Segoe UI", 10),
    )
    style.configure(
        "ComparePrimary.TButton",
        background="#207C7E",
        foreground="white",
        padding=(18, 10),
        font=("Segoe UI", 10, "bold"),
    )
    style.map(
        "ComparePrimary.TButton",
        background=[
            ("disabled", "#B2C3CA"),
            ("active", "#176769"),
        ],
    )
    style.configure(
        "Compare.TCheckbutton",
        background="#F4F7FB",
        font=("Segoe UI", 10),
    )

    # При отдельном запуске основное окно ещё не настроило календарь.
    if parent is None:
        for name in (
            "CalendarDay",
            "Weekend",
            "Today",
            "SelectedDay",
            "CalendarNav",
        ):
            style.configure(
                f"{name}.TButton",
                padding=5,
                font=("Segoe UI", 10),
            )

        style.configure(
            "SelectedDay.TButton",
            background="#207C7E",
            foreground="white",
        )
        style.configure(
            "Today.TButton",
            foreground="#207C7E",
        )
        style.configure(
            "Weekend.TButton",
            foreground="#A34F68",
        )
        style.configure(
            "CalendarTitle.TLabel",
            font=("Segoe UI", 11, "bold"),
        )

    today = date.today()
    previous = today.replace(day=1) - timedelta(days=1)

    defaults = {
        "previous": "",
        "current": "",
        "day_previous": "",
        "day_current": "",
        "template": str(TEMPLATE),
        "from0": previous.replace(day=7).strftime("%d.%m.%Y"),
        "to0": previous.replace(day=20).strftime("%d.%m.%Y"),
        "from1": today.replace(day=7).strftime("%d.%m.%Y"),
        "to1": today.replace(day=20).strftime("%d.%m.%Y"),
    }

    values = {
        key: tk.StringVar(master=window, value=value)
        for key, value in defaults.items()
    }

    complete = tk.BooleanVar(master=window)
    hint = tk.StringVar(master=window)
    status = tk.StringVar(
        master=window,
        value="Выберите две выгрузки.",
    )

    messages = queue.Queue()
    state = {"busy": False, "result": None}
    controls = []

    outer = ttk.Frame(
        window,
        padding=14,
        style="Compare.TFrame",
    )
    outer.pack(fill="both", expand=True)
    outer.columnconfigure(0, weight=1)
    outer.rowconfigure(0, weight=1)

    # Прокручиваемая форма. Нижняя панель остаётся на месте.
    canvas = tk.Canvas(
        outer,
        background="#F4F7FB",
        highlightthickness=0,
    )
    scroll = ttk.Scrollbar(
        outer,
        orient="vertical",
        command=canvas.yview,
    )
    canvas.configure(yscrollcommand=scroll.set)

    canvas.grid(row=0, column=0, sticky="nsew")
    scroll.grid(row=0, column=1, sticky="ns")

    form = ttk.Frame(
        canvas,
        style="Compare.TFrame",
        padding=(0, 0, 10, 0),
    )
    form.columnconfigure(0, weight=1)

    form_id = canvas.create_window(
        (0, 0),
        window=form,
        anchor="nw",
    )

    form.bind(
        "<Configure>",
        lambda event: canvas.configure(
            scrollregion=canvas.bbox("all")
        ),
    )
    canvas.bind(
        "<Configure>",
        lambda event: canvas.itemconfigure(
            form_id,
            width=event.width,
        ),
    )

    ttk.Label(
        form,
        text="Сравнение периодов",
        style="CompareTitle.TLabel",
    ).grid(row=0, column=0, sticky="w")

    ttk.Label(
        form,
        text="Выберите файлы и периоды для сравнения",
        style="CompareHint.TLabel",
    ).grid(
        row=1,
        column=0,
        sticky="w",
        pady=(4, 14),
    )

    def card(row, title):
        box = ttk.Frame(
            form,
            padding=14,
            style="CompareCard.TFrame",
        )
        box.grid(
            row=row,
            column=0,
            sticky="ew",
            pady=(0, 12),
        )
        box.columnconfigure(1, weight=1)

        ttk.Label(
            box,
            text=title,
            style="CompareHeading.TLabel",
        ).grid(
            row=0,
            column=0,
            columnspan=3,
            sticky="w",
            pady=(0, 8),
        )

        return box

    files = card(2, "1. Исходные файлы")

    def choose_file(key):
        patterns = (
            "*.xlsx"
            if key == "template"
            else "*.xls *.xlsx *.xlsm"
        )

        path = filedialog.askopenfilename(
            parent=window,
            title="Выберите файл",
            filetypes=[("Excel", patterns)],
        )

        if path:
            values[key].set(path)

    for row, (key, title) in enumerate(
        (
            ("previous", "Предыдущий месяц"),
            ("day_previous", "Стационар — предыдущий"),
            ("current", "Текущий месяц"),
            ("day_current", "Стационар — текущий"),
            ("template", "Шаблон отчёта"),
        ),
        1,
    ):
        ttk.Label(
            files,
            text=title,
            style="Compare.TLabel",
        ).grid(
            row=row,
            column=0,
            sticky="w",
            pady=5,
        )

        entry = ttk.Entry(
            files,
            textvariable=values[key],
            state="readonly",
            style="Compare.TEntry",
            font=("Segoe UI", 10),
        )
        entry.grid(
            row=row,
            column=1,
            sticky="ew",
            padx=10,
            pady=5,
        )

        browse = ttk.Button(
            files,
            text="Выбрать…",
            style="Compare.TButton",
            command=lambda k=key: choose_file(k),
        )
        browse.grid(row=row, column=2, pady=5)

        controls.extend((entry, browse))
        if key.startswith("day_"):
            clear = ttk.Button(
                files, text="Очистить", style="Compare.TButton",
                command=lambda k=key: values[k].set(""),
            )
            clear.grid(row=row, column=3, padx=(6, 0), pady=5)
            controls.append(clear)

    ttk.Label(
        files, style="CompareHint.TLabel", wraplength=820,
        text=("Стационар необязателен: выберите оба файла или оставьте оба поля пустыми. "
              "Периоды ниже применяются к услугам ЖК и к дате выбытия из стационара."),
    ).grid(row=6, column=0, columnspan=4, sticky="w", pady=(4, 0))

    periods = card(3, "2. Периоды сравнения")
    periods.columnconfigure(0, weight=1, uniform="period")
    periods.columnconfigure(1, weight=1, uniform="period")

    def set_period(index, start, end):
        values[f"from{index}"].set(
            start.strftime("%d.%m.%Y")
        )
        values[f"to{index}"].set(
            end.strftime("%d.%m.%Y")
        )

    def open_calendar(index, edge, anchor):
        key = (
            f"from{index}"
            if edge == "start"
            else f"to{index}"
        )

        def selected(chosen):
            values[key].set(chosen.strftime("%d.%m.%Y"))

        CalendarPopup(
            window,
            initial_date=parse_date(values[key].get()),
            on_select=selected,
            anchor_widget=anchor,
        )

    for index, title in enumerate(
        ("Предыдущий месяц", "Текущий месяц")
    ):
        box = ttk.Frame(
            periods,
            style="CompareCard.TFrame",
        )
        box.grid(
            row=1,
            column=index,
            sticky="ew",
            padx=(0, 16) if index == 0 else 0,
        )
        box.columnconfigure(1, weight=1)

        ttk.Label(
            box,
            text=title,
            style="CompareHeading.TLabel",
        ).grid(
            row=0,
            column=0,
            columnspan=3,
            sticky="w",
            pady=(4, 8),
        )

        for row, (edge, caption, key) in enumerate(
            (
                ("start", "С", f"from{index}"),
                ("end", "По", f"to{index}"),
            ),
            1,
        ):
            ttk.Label(
                box,
                text=caption,
                style="Compare.TLabel",
            ).grid(
                row=row,
                column=0,
                sticky="w",
                padx=(0, 8),
            )

            entry = ttk.Entry(
                box,
                textvariable=values[key],
                state="readonly",
                width=12,
                font=("Segoe UI", 11),
                style="Compare.TEntry",
            )
            entry.grid(
                row=row,
                column=1,
                sticky="ew",
                pady=4,
            )

            picker = ttk.Button(
                box,
                text="Календарь",
                style="Compare.TButton",
            )
            picker.configure(
                command=lambda i=index, e=edge, b=picker: (
                    open_calendar(i, e, b)
                )
            )
            picker.grid(
                row=row,
                column=2,
                padx=(8, 0),
                pady=4,
            )

            def activate(
                event,
                i=index,
                e=edge,
                anchor=entry,
            ):
                if not state["busy"]:
                    open_calendar(i, e, anchor)

                return "break"

            entry.bind("<Button-1>", activate)
            entry.bind("<Return>", activate)
            entry.bind("<space>", activate)

            controls.extend((entry, picker))

    quick = ttk.Frame(
        periods,
        style="CompareCard.TFrame",
    )
    quick.grid(
        row=2,
        column=0,
        columnspan=2,
        sticky="w",
        pady=(10, 0),
    )

    ttk.Label(
        quick,
        text="Быстрый выбор:",
        style="Compare.TLabel",
    ).pack(side="left", padx=(0, 8))

    def preset(day):
        for index in (0, 1):
            month = parse_date(
                values[f"from{index}"].get()
            )
            start, end = _comparison_span(
                month.replace(day=day)
            )
            set_period(index, start, end)

    for day in (1, 7, 15):
        btn = ttk.Button(
            quick,
            text=f"{day:02d}–{day + 13:02d}",
            style="Compare.TButton",
            command=lambda d=day: preset(d),
        )
        btn.pack(side="left", padx=(0, 6))
        controls.append(btn)

    def align_previous():
        start = parse_date(values["from1"].get())
        end = parse_date(values["to1"].get())

        if start > end or (start.year, start.month) != (end.year, end.month):
            messagebox.showinfo(
                "Проверьте даты",
                "Сначала выберите корректный период текущего месяца.",
                parent=window,
            )
            return

        previous_end = start.replace(day=1) - timedelta(days=1)
        a = previous_end.replace(
            day=min(start.day, previous_end.day)
        )

        if end.day == monthrange(end.year, end.month)[1]:
            b = previous_end
        else:
            b = previous_end.replace(
                day=min(end.day, previous_end.day)
            )

        set_period(0, a, b)

    align = ttk.Button(
        periods,
        text="Такие же дни в предыдущем месяце",
        style="Compare.TButton",
        command=align_previous,
    )
    align.grid(
        row=3,
        column=0,
        columnspan=2,
        sticky="w",
        pady=(8, 0),
    )
    controls.append(align)

    hint_label = ttk.Label(
        form,
        textvariable=hint,
        style="CompareHint.TLabel",
    )
    hint_label.grid(
        row=4,
        column=0,
        sticky="w",
        pady=(0, 8),
    )

    check = ttk.Checkbutton(
        form,
        variable=complete,
        style="Compare.TCheckbutton",
        text=(
            "Все выбранные выгрузки полные за выбранные периоды; "
            "дни без услуг допустимы"
        ),
    )
    check.grid(
        row=5,
        column=0,
        sticky="w",
        pady=(0, 12),
    )
    controls.append(check)

    log_box = card(6, "Результат и сообщения")
    log_box.rowconfigure(1, weight=1)
    log_box.columnconfigure(0, weight=1)

    log = tk.Text(
        log_box,
        height=6,
        wrap="word",
        font=("Segoe UI", 10),
        background="white",
        foreground="#243B53",
        relief="flat",
        state="disabled",
    )
    log.grid(
        row=1,
        column=0,
        columnspan=2,
        sticky="nsew",
    )

    log_scroll = ttk.Scrollbar(
        log_box,
        command=log.yview,
    )
    log_scroll.grid(
        row=1,
        column=2,
        sticky="ns",
    )
    log.configure(yscrollcommand=log_scroll.set)

    def set_log(text):
        log.configure(state="normal")
        log.delete("1.0", "end")
        log.insert("end", text)
        log.configure(state="disabled")

    set_log(
        "Здесь появятся результаты проверки "
        "и путь к готовому отчёту."
    )

    footer = ttk.Frame(
        outer,
        padding=(0, 12, 0, 0),
        style="Compare.TFrame",
    )
    footer.grid(
        row=1,
        column=0,
        columnspan=2,
        sticky="ew",
    )
    footer.columnconfigure(0, weight=1)

    ttk.Label(
        footer,
        textvariable=status,
        style="CompareHint.TLabel",
    ).grid(
        row=0,
        column=0,
        columnspan=3,
        sticky="w",
        pady=(0, 8),
    )

    progress = ttk.Progressbar(
        footer,
        mode="indeterminate",
    )
    progress.grid(
        row=1,
        column=0,
        sticky="ew",
        padx=(0, 12),
    )

    save_button = ttk.Button(
        footer,
        text="Сохранить отчёт…",
        style="Compare.TButton",
    )
    save_button.grid(
        row=1,
        column=1,
        padx=(0, 8),
    )

    button = ttk.Button(
        footer,
        text="Сформировать отчёт",
        style="ComparePrimary.TButton",
    )
    button.grid(row=1, column=2)

    def read_dates():
        return tuple(
            parse_date(values[key].get())
            for key in (
                "from0", "to0",
                "from1", "to1",
            )
        )

    def refresh():
        valid = True

        try:
            a, b, c, d = read_dates()
            validate_periods(a, b, c, d)

            hint.set(
                f"Предыдущий период: {(b - a).days + 1} дн.; "
                f"текущий: {(d - c).days + 1} дн. "
                "Обе даты включены. Начало и конец выбираются отдельно."
            )
            hint_label.configure(
                foreground="#526779"
            )

        except (ValueError, ReportError) as exc:
            valid = False
            hint.set(str(exc))
            hint_label.configure(
                foreground="#B42318"
            )

        paired_hospital_files = bool(values["day_previous"].get()) == bool(values["day_current"].get())
        if valid and not paired_hospital_files:
            hint.set("Выберите оба файла стационара или очистите оба поля стационара.")
            hint_label.configure(foreground="#B42318")

        ready = (
            valid
            and paired_hospital_files
            and complete.get()
            and all(
                values[key].get()
                for key in (
                    "previous",
                    "current",
                    "template",
                )
            )
        )

        button.state(
            ["!disabled"]
            if ready and not state["busy"]
            else ["disabled"]
        )

        save_button.state(
            ["!disabled"]
            if state["result"] and not state["busy"]
            else ["disabled"]
        )

    def changed(*_):
        if state["result"] is not None:
            state["result"] = None

            status.set(
                "Параметры изменены. "
                "Сформируйте отчёт заново."
            )
            set_log(
                "Параметры изменены. Предыдущий результат "
                "больше не выбран для сохранения."
            )

        refresh()

    for var in values.values():
        var.trace_add("write", changed)

    complete.trace_add(
        "write",
        lambda *_: refresh(),
    )

    def set_busy(busy):
        state["busy"] = busy

        for widget in controls:
            widget.state(
                ["disabled"]
                if busy
                else ["!disabled"]
            )

        if busy:
            progress.start(12)
        else:
            progress.stop()

        refresh()

    def save_report():
        if not state["result"]:
            return

        content, report, sources, dates = state["result"]

        output = filedialog.asksaveasfilename(
            parent=window,
            defaultextension=".xlsx",
            initialfile=(
                f"Анализ_периодов_"
                f"{dates[0]:%Y-%m-%d}_"
                f"{dates[2]:%Y-%m-%d}.xlsx"
            ),
            filetypes=[("Excel", "*.xlsx")],
        )

        if not output:
            status.set(
                "Отчёт готов. Его можно сохранить "
                "кнопкой «Сохранить отчёт…»."
            )
            return

        target = Path(output)

        try:
            if target.suffix.lower() != ".xlsx":
                raise ReportError(
                    "Расширение результата должно быть .xlsx."
                )

            if target.resolve() in {
                path.resolve()
                for path in sources
            }:
                raise ReportError(
                    "Выберите другое имя: "
                    "это исходный файл или шаблон."
                )

            with target.open("xb") as handle:
                handle.write(content)

        except FileExistsError:
            messagebox.showerror(
                "Файл уже существует",
                "Нажмите «Сохранить отчёт…» "
                "и выберите новое имя.",
                parent=window,
            )

        except Exception as exc:
            messagebox.showerror(
                "Ошибка сохранения",
                str(exc),
                parent=window,
            )

        else:
            status.set("Готово. Отчёт сохранён.")
            set_log(
                report
                + f"\n\nСохранён файл:\n{target}"
            )

    def poll():
        try:
            ok, result = messages.get_nowait()
        except queue.Empty:
            window.after(100, poll)
            return

        if ok:
            state["result"] = result
            set_log(result[1])
            status.set("Расчёт завершён.")
        else:
            set_log(str(result))
            status.set(
                "Не удалось сформировать отчёт."
            )

        set_busy(False)

        if ok:
            save_report()
        else:
            messagebox.showerror(
                "Ошибка",
                str(result),
                parent=window,
            )

    def start():
        if state["busy"]:
            return

        try:
            paths = [
                Path(values[key].get())
                for key in (
                    "previous",
                    "current",
                    "template",
                )
            ]
            hospital_paths = [
                Path(values[key].get().strip()) if values[key].get().strip() else None
                for key in ("day_previous", "day_current")
            ]

            dates = read_dates()
            validate_periods(*dates)

            if not complete.get():
                raise ReportError(
                    "Подтвердите полноту всех выбранных выгрузок."
                )

            if not all(path.is_file() for path in paths):
                raise ReportError(
                    "Выберите существующие выгрузки "
                    "и шаблон."
                )

            if paths[0].resolve() == paths[1].resolve():
                raise ReportError(
                    "Выбран один и тот же "
                    "исходный файл дважды."
                )
            validate_hospital_files(hospital_paths, paths)

        except Exception as exc:
            messagebox.showerror(
                "Проверка данных",
                str(exc),
                parent=window,
            )
            return

        state["result"] = None
        set_busy(True)
        status.set("Чтение файлов и расчёт…")
        set_log(
            "Обработка выполняется. "
            "Дождитесь завершения расчёта."
        )

        def work():
            try:
                hash0 = hashlib.sha256(
                    paths[0].read_bytes()
                ).digest()
                hash1 = hashlib.sha256(
                    paths[1].read_bytes()
                ).digest()

                if hash0 == hash1:
                    raise ReportError(
                        "Содержимое двух исходных "
                        "файлов совпадает."
                    )

                old = read_period(
                    paths[0],
                    dates[0],
                    dates[1],
                    day_hospital_path=hospital_paths[0],
                )
                new = read_period(
                    paths[1],
                    dates[2],
                    dates[3],
                    day_hospital_path=hospital_paths[1],
                )

                book, unknown = build_book(
                    paths[2],
                    old,
                    new,
                    dates,
                )

                report = (
                    describe(
                        old,
                        "Предыдущий период",
                    )
                    + "\n\n"
                    + describe(
                        new,
                        "Текущий период",
                        unknown,
                    )
                )

                stream = BytesIO()

                try:
                    book.save(stream)
                finally:
                    book.close()

                messages.put((
                    True,
                    (
                        stream.getvalue(),
                        report,
                        paths + [path for path in hospital_paths if path is not None],
                        dates,
                    ),
                ))

            except Exception as exc:
                messages.put((False, str(exc)))

        threading.Thread(
            target=work,
            daemon=True,
        ).start()

        window.after(100, poll)

    button.configure(command=start)
    save_button.configure(command=save_report)

    def wheel(event):
        if event.widget is not log:
            step = -1 if event.delta > 0 else 1
            canvas.yview_scroll(
                step * 3,
                "units",
            )

    window.bind(
        "<MouseWheel>",
        wheel,
        add="+",
    )

    def close_window():
        if state["busy"]:
            messagebox.showinfo(
                "Обработка",
                "Дождитесь завершения расчёта.",
                parent=window,
            )
        else:
            window.destroy()

    window.protocol(
        "WM_DELETE_WINDOW",
        close_window,
    )

    refresh()
    window.update_idletasks()

    x = max(
        0,
        (window.winfo_screenwidth() - width) // 2,
    )
    y = max(
        0,
        (window.winfo_screenheight() - height) // 2,
    )

    window.geometry(
        f"{width}x{height}+{x}+{y}"
    )
    window.deiconify()

    if parent is None:
        window.mainloop()

if __name__ == "__main__":
    launch_compare()
