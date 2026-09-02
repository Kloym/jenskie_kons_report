from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
from datetime import date, timedelta
from decimal import Decimal, ROUND_HALF_UP
from typing import Mapping, Protocol, Sequence

from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


MONEY_QUANT = Decimal("0.01")


class ServiceRecord(Protocol):
    """Минимальный интерфейс записи из основной программы."""

    doctor: str
    service_code: str
    service_name: str
    tariff: Decimal
    quantity: Decimal


@dataclass(frozen=True)
class AnalyticsTheme:
    """Палитра аналитики, согласованная с основным отчётом."""

    primary: str
    accent: str
    pale_primary: str
    pale_accent: str
    text: str = "203040"
    muted: str = "63788A"
    grid: str = "D5E0E8"
    white: str = "FFFFFF"


@dataclass(frozen=True)
class ServiceSummary:
    code: str
    name: str
    quantity: Decimal
    amount: Decimal


@dataclass(frozen=True)
class DoctorSummary:
    doctor: str
    quantity: Decimal
    amount: Decimal


@dataclass(frozen=True)
class TimePoint:
    label: str
    quantity: Decimal
    amount: Decimal


@dataclass(frozen=True)
class AnalyticsData:
    doctor_count: int
    unique_service_count: int
    total_quantity: Decimal
    total_amount: Decimal
    top_services: tuple[ServiceSummary, ...]
    top_doctors: tuple[DoctorSummary, ...]
    time_points: tuple[TimePoint, ...]
    time_granularity: str


class AnalyticsBuilder:
    """Создаёт лист «Аналитика» с KPI, таблицами и диаграммами."""

    SHEET_NAME = "Аналитика"
    TOP_ITEMS = 10
    DAILY_TREND_LIMIT = 62

    def __init__(self, theme: AnalyticsTheme) -> None:
        self.theme = theme
        self._thin_side = Side(style="thin", color=theme.grid)

    def build(
        self,
        workbook: object,
        department_name: str,
        records: Sequence[ServiceRecord],
        daily_quantity: Mapping[date, Decimal],
        daily_amount: Mapping[date, Decimal],
        date_from: date,
        date_to: date,
    ) -> object:
        """Добавляет или полностью пересоздаёт аналитический лист."""
        if self.SHEET_NAME in workbook.sheetnames:
            del workbook[self.SHEET_NAME]

        data = self._prepare_data(
            records,
            daily_quantity,
            daily_amount,
            date_from,
            date_to,
        )

        worksheet = workbook.create_sheet(self.SHEET_NAME)
        self._configure_sheet(worksheet)
        self._write_header(
            worksheet,
            department_name,
            date_from,
            date_to,
        )
        self._write_kpi_cards(worksheet, data)
        self._write_services_section(worksheet, data.top_services)
        self._write_doctors_section(worksheet, data.top_doctors)
        last_row = self._write_trend_section(
            worksheet,
            data.time_points,
            data.time_granularity,
        )
        self._finish_layout(worksheet, last_row)

        # При открытии книги пользователь сначала видит подробный отчёт.
        workbook.active = 0
        return worksheet

    def _prepare_data(
        self,
        records: Sequence[ServiceRecord],
        daily_quantity: Mapping[date, Decimal],
        daily_amount: Mapping[date, Decimal],
        date_from: date,
        date_to: date,
    ) -> AnalyticsData:
        services: dict[tuple[str, str], list[Decimal]] = defaultdict(
            lambda: [Decimal("0"), Decimal("0")]
        )
        doctors: dict[str, list[Decimal]] = defaultdict(
            lambda: [Decimal("0"), Decimal("0")]
        )

        total_quantity = Decimal("0")
        total_amount = Decimal("0")

        for record in records:
            quantity = _as_decimal(record.quantity)
            tariff = _as_decimal(record.tariff)
            amount = (tariff * quantity).quantize(
                MONEY_QUANT,
                rounding=ROUND_HALF_UP,
            )

            service_key = (str(record.service_code), str(record.service_name))
            services[service_key][0] += quantity
            services[service_key][1] += amount

            doctor = str(record.doctor)
            doctors[doctor][0] += quantity
            doctors[doctor][1] += amount

            total_quantity += quantity
            total_amount += amount

        service_rows = sorted(
            (
                ServiceSummary(code, name, values[0], values[1])
                for (code, name), values in services.items()
            ),
            key=lambda item: (-item.quantity, item.code.casefold(), item.name.casefold()),
        )
        doctor_rows = sorted(
            (
                DoctorSummary(doctor, values[0], values[1])
                for doctor, values in doctors.items()
            ),
            key=lambda item: (-item.quantity, item.doctor.casefold()),
        )

        time_points, granularity = self._build_time_points(
            daily_quantity,
            daily_amount,
            date_from,
            date_to,
        )

        return AnalyticsData(
            doctor_count=len(doctors),
            unique_service_count=len(services),
            total_quantity=total_quantity,
            total_amount=total_amount.quantize(MONEY_QUANT, rounding=ROUND_HALF_UP),
            top_services=tuple(service_rows[: self.TOP_ITEMS]),
            top_doctors=tuple(doctor_rows[: self.TOP_ITEMS]),
            time_points=time_points,
            time_granularity=granularity,
        )

    def _build_time_points(
        self,
        daily_quantity: Mapping[date, Decimal],
        daily_amount: Mapping[date, Decimal],
        date_from: date,
        date_to: date,
    ) -> tuple[tuple[TimePoint, ...], str]:
        period_days = (date_to - date_from).days + 1
        if period_days <= 1:
            return (), "day"

        if period_days <= self.DAILY_TREND_LIMIT:
            points: list[TimePoint] = []
            current_date = date_from
            while current_date <= date_to:
                points.append(
                    TimePoint(
                        label=current_date.strftime("%d.%m"),
                        quantity=_as_decimal(
                            daily_quantity.get(current_date, Decimal("0"))
                        ),
                        amount=_as_decimal(
                            daily_amount.get(current_date, Decimal("0"))
                        ).quantize(MONEY_QUANT, rounding=ROUND_HALF_UP),
                    )
                )
                current_date += timedelta(days=1)
            return tuple(points), "day"

        # Для длинных периодов дневной график становится нечитаемым.
        monthly_quantity: dict[tuple[int, int], Decimal] = defaultdict(
            lambda: Decimal("0")
        )
        monthly_amount: dict[tuple[int, int], Decimal] = defaultdict(
            lambda: Decimal("0")
        )
        for service_date, value in daily_quantity.items():
            monthly_quantity[(service_date.year, service_date.month)] += _as_decimal(value)
        for service_date, value in daily_amount.items():
            monthly_amount[(service_date.year, service_date.month)] += _as_decimal(value)

        points = []
        current_month = date(date_from.year, date_from.month, 1)
        last_month = date(date_to.year, date_to.month, 1)
        while current_month <= last_month:
            key = (current_month.year, current_month.month)
            points.append(
                TimePoint(
                    label=current_month.strftime("%m.%Y"),
                    quantity=monthly_quantity[key],
                    amount=monthly_amount[key].quantize(
                        MONEY_QUANT,
                        rounding=ROUND_HALF_UP,
                    ),
                )
            )
            if current_month.month == 12:
                current_month = date(current_month.year + 1, 1, 1)
            else:
                current_month = date(
                    current_month.year,
                    current_month.month + 1,
                    1,
                )
        return tuple(points), "month"

    def _configure_sheet(self, worksheet: object) -> None:
        worksheet.sheet_view.showGridLines = False
        worksheet.sheet_view.zoomScale = 90
        worksheet.freeze_panes = "A10"
        worksheet.sheet_properties.tabColor = self.theme.accent

        widths = {
            "A": 14,
            "B": 43,
            "C": 14,
            "D": 18,
            "E": 3,
            "F": 13,
            "G": 13,
            "H": 13,
            "I": 13,
            "J": 13,
            "K": 13,
            "L": 13,
            "M": 13,
            "N": 13,
        }
        for column, width in widths.items():
            worksheet.column_dimensions[column].width = width

    def _write_header(
        self,
        worksheet: object,
        department_name: str,
        date_from: date,
        date_to: date,
    ) -> None:
        self._style_block(
            worksheet,
            min_row=1,
            max_row=2,
            min_col=1,
            max_col=14,
            fill=self.theme.primary,
        )
        worksheet.merge_cells("A1:N2")
        cell = worksheet["A1"]
        cell.value = "АНАЛИТИКА ПО ОКАЗАННЫМ УСЛУГАМ"
        cell.font = Font(
            name="Aptos Display",
            size=18,
            bold=True,
            color=self.theme.white,
        )
        cell.alignment = Alignment(horizontal="center", vertical="center")
        worksheet.row_dimensions[1].height = 25
        worksheet.row_dimensions[2].height = 25

        self._style_block(
            worksheet,
            min_row=3,
            max_row=3,
            min_col=1,
            max_col=14,
            fill=self.theme.accent,
        )
        worksheet.merge_cells("A3:N3")
        subtitle = worksheet["A3"]
        subtitle.value = (
            f"{department_name}  •  "
            f"период {date_from:%d.%m.%Y} — {date_to:%d.%m.%Y}"
        )
        subtitle.font = Font(
            name="Aptos",
            size=11,
            bold=True,
            color=self.theme.white,
        )
        subtitle.alignment = Alignment(horizontal="center", vertical="center")
        worksheet.row_dimensions[3].height = 25
        worksheet.row_dimensions[4].height = 8

    def _write_kpi_cards(self, worksheet: object, data: AnalyticsData) -> None:
        cards = (
            ("A5:C5", "A6:C7", "Врачей", data.doctor_count, "#,##0"),
            (
                "D5:F5",
                "D6:F7",
                "Оказано услуг",
                _excel_number(data.total_quantity),
                "#,##0",
            ),
            (
                "G5:I5",
                "G6:I7",
                "Уникальных услуг",
                data.unique_service_count,
                "#,##0",
            ),
            (
                "J5:N5",
                "J6:N7",
                "Стоимость оказанных услуг, руб.",
                float(data.total_amount),
                "#,##0.00",
            ),
        )
        for label_range, value_range, label, value, number_format in cards:
            self._write_card(
                worksheet,
                label_range,
                value_range,
                label,
                value,
                number_format,
            )
        worksheet.row_dimensions[5].height = 22
        worksheet.row_dimensions[6].height = 25
        worksheet.row_dimensions[7].height = 25
        worksheet.row_dimensions[8].height = 8

    def _write_card(
        self,
        worksheet: object,
        label_range: str,
        value_range: str,
        label: str,
        value: object,
        number_format: str,
    ) -> None:
        label_cells = worksheet[label_range]
        value_cells = worksheet[value_range]
        for row in label_cells:
            for cell in row:
                cell.fill = PatternFill("solid", fgColor=self.theme.pale_primary)
                cell.border = Border(top=self._thin_side, left=self._thin_side, right=self._thin_side)
        for row in value_cells:
            for cell in row:
                cell.fill = PatternFill("solid", fgColor=self.theme.white)
                cell.border = Border(bottom=self._thin_side, left=self._thin_side, right=self._thin_side)

        worksheet.merge_cells(label_range)
        worksheet.merge_cells(value_range)
        label_cell = worksheet[label_range.split(":")[0]]
        value_cell = worksheet[value_range.split(":")[0]]

        label_cell.value = label
        label_cell.font = Font(
            name="Aptos",
            size=9,
            bold=True,
            color=self.theme.muted,
        )
        label_cell.alignment = Alignment(horizontal="center", vertical="center")

        value_cell.value = value
        value_cell.font = Font(
            name="Aptos Display",
            size=18,
            bold=True,
            color=self.theme.primary,
        )
        value_cell.alignment = Alignment(horizontal="center", vertical="center")
        value_cell.number_format = number_format

    def _write_services_section(
        self,
        worksheet: object,
        services: Sequence[ServiceSummary],
    ) -> None:
        title_row = 9
        header_row = 10
        first_data_row = 11
        self._write_section_title(
            worksheet,
            title_row,
            "Топ-10 услуг по количеству",
            end_column=4,
        )
        self._write_table_header(
            worksheet,
            header_row,
            ("Код", "Наименование услуги", "Кол-во", "Сумма, руб."),
        )

        for offset, item in enumerate(services):
            row = first_data_row + offset
            values = (
                str(item.code),
                str(item.name),
                _excel_number(item.quantity),
                float(item.amount),
            )
            self._write_data_row(worksheet, row, values, offset)
            worksheet.cell(row, 1).number_format = "@"
            worksheet.cell(row, 3).number_format = "#,##0"
            worksheet.cell(row, 4).number_format = "#,##0.00"
            worksheet.row_dimensions[row].height = 30

        if services:
            last_data_row = first_data_row + len(services) - 1
            chart = self._make_bar_chart(
                worksheet,
                title="Какие услуги оказываются чаще всего",
                header_row=header_row,
                first_data_row=first_data_row,
                last_data_row=last_data_row,
                category_column=1,
                value_column=3,
                show_category_in_labels=True,
            )
            worksheet.add_chart(chart, "F11")

    def _write_doctors_section(
        self,
        worksheet: object,
        doctors: Sequence[DoctorSummary],
    ) -> None:
        title_row = 29
        header_row = 30
        first_data_row = 31
        self._write_section_title(
            worksheet,
            title_row,
            "Топ-10 врачей по количеству услуг",
            end_column=4,
        )
        self._write_table_header(
            worksheet,
            header_row,
            ("Врач", "", "Кол-во", "Сумма, руб."),
        )
        worksheet.merge_cells(start_row=30, start_column=1, end_row=30, end_column=2)

        for offset, item in enumerate(doctors):
            row = first_data_row + offset
            worksheet.merge_cells(
                start_row=row,
                start_column=1,
                end_row=row,
                end_column=2,
            )
            values = (
                str(item.doctor),
                None,
                _excel_number(item.quantity),
                float(item.amount),
            )
            self._write_data_row(worksheet, row, values, offset)
            worksheet.cell(row, 3).number_format = "#,##0"
            worksheet.cell(row, 4).number_format = "#,##0.00"
            worksheet.row_dimensions[row].height = 24

        if doctors:
            last_data_row = first_data_row + len(doctors) - 1
            chart = self._make_bar_chart(
                worksheet,
                title="Распределение нагрузки по врачам",
                header_row=header_row,
                first_data_row=first_data_row,
                last_data_row=last_data_row,
                category_column=1,
                value_column=3,
                show_category_in_labels=False,
            )
            worksheet.add_chart(chart, "F29")

    def _write_trend_section(
        self,
        worksheet: object,
        points: Sequence[TimePoint],
        granularity: str,
    ) -> int:
        title_row = 49
        header_row = 50
        first_data_row = 51
        period_word = "дням" if granularity == "day" else "месяцам"
        self._write_section_title(
            worksheet,
            title_row,
            f"Динамика по {period_word}",
            end_column=4,
        )

        if not points:
            self._style_block(
                worksheet,
                min_row=50,
                max_row=55,
                min_col=1,
                max_col=14,
                fill=self.theme.pale_accent,
            )
            worksheet.merge_cells("A50:N55")
            message = worksheet["A50"]
            message.value = (
                "Тут возможна динамика, но она не строится для одного дня. =) "
            )
            message.font = Font(
                name="Aptos",
                size=12,
                bold=True,
                color=self.theme.muted,
            )
            message.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )
            return 55

        self._write_table_header(
            worksheet,
            header_row,
            ("Период", "", "Кол-во", "Сумма, руб."),
        )
        worksheet.merge_cells(start_row=50, start_column=1, end_row=50, end_column=2)

        for offset, point in enumerate(points):
            row = first_data_row + offset
            worksheet.merge_cells(
                start_row=row,
                start_column=1,
                end_row=row,
                end_column=2,
            )
            values = (
                point.label,
                None,
                _excel_number(point.quantity),
                float(point.amount),
            )
            self._write_data_row(worksheet, row, values, offset)
            worksheet.cell(row, 3).number_format = "#,##0"
            worksheet.cell(row, 4).number_format = "#,##0.00"

        last_data_row = first_data_row + len(points) - 1
        chart = self._make_line_chart(
            worksheet,
            title=f"Количество услуг по {period_word}",
            header_row=header_row,
            first_data_row=first_data_row,
            last_data_row=last_data_row,
            category_column=1,
            value_column=3,
        )
        worksheet.add_chart(chart, "F49")
        return max(last_data_row, 66)

    def _write_section_title(
        self,
        worksheet: object,
        row: int,
        title: str,
        end_column: int,
    ) -> None:
        self._style_block(
            worksheet,
            min_row=row,
            max_row=row,
            min_col=1,
            max_col=end_column,
            fill=self.theme.accent,
        )
        worksheet.merge_cells(
            start_row=row,
            start_column=1,
            end_row=row,
            end_column=end_column,
        )
        cell = worksheet.cell(row, 1)
        cell.value = title
        cell.font = Font(
            name="Aptos",
            size=10,
            bold=True,
            color=self.theme.white,
        )
        cell.alignment = Alignment(horizontal="left", vertical="center")
        worksheet.row_dimensions[row].height = 25

    def _write_table_header(
        self,
        worksheet: object,
        row: int,
        headers: Sequence[str],
    ) -> None:
        for column, header in enumerate(headers, start=1):
            cell = worksheet.cell(row, column, value=header)
            cell.fill = PatternFill("solid", fgColor=self.theme.primary)
            cell.font = Font(
                name="Aptos",
                size=9,
                bold=True,
                color=self.theme.white,
            )
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )
            cell.border = Border(bottom=self._thin_side)
        worksheet.row_dimensions[row].height = 25

    def _write_data_row(
        self,
        worksheet: object,
        row: int,
        values: Sequence[object],
        offset: int,
    ) -> None:
        fill_color = (
            self.theme.white if offset % 2 == 0 else self.theme.pale_accent
        )
        for column, value in enumerate(values, start=1):
            cell = worksheet.cell(row, column, value=value)
            cell.fill = PatternFill("solid", fgColor=fill_color)
            cell.font = Font(name="Aptos", size=9, color=self.theme.text)
            cell.border = Border(bottom=self._thin_side)
            cell.alignment = Alignment(
                horizontal="left" if column <= 2 else "right",
                vertical="center",
                wrap_text=column == 2,
            )

    def _make_bar_chart(
        self,
        worksheet: object,
        title: str,
        header_row: int,
        first_data_row: int,
        last_data_row: int,
        category_column: int,
        value_column: int,
        show_category_in_labels: bool,
    ) -> BarChart:
        chart = BarChart()
        chart.type = "bar"
        chart.style = 10
        chart.title = title
        chart.height = 8.6
        chart.width = 19.0
        chart.legend = None
        chart.gapWidth = 55
        chart.varyColors = False
        chart.x_axis.numFmt = "#,##0"
        chart.x_axis.title = None
        chart.y_axis.title = None

        data = Reference(
            worksheet,
            min_col=value_column,
            min_row=header_row,
            max_row=last_data_row,
        )
        categories = Reference(
            worksheet,
            min_col=category_column,
            min_row=first_data_row,
            max_row=last_data_row,
        )
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)

        labels = DataLabelList()
        labels.showVal = True
        labels.showLegendKey = False
        labels.showCatName = show_category_in_labels
        labels.showSerName = False
        if show_category_in_labels:
            labels.separator = " • Кол-во: "
        chart.dLbls = labels
        self._style_chart_series(chart)
        return chart

    def _make_line_chart(
        self,
        worksheet: object,
        title: str,
        header_row: int,
        first_data_row: int,
        last_data_row: int,
        category_column: int,
        value_column: int,
    ) -> LineChart:
        chart = LineChart()
        chart.style = 13
        chart.title = title
        chart.height = 7.3
        chart.width = 17.0
        chart.legend = None
        chart.y_axis.title = "Количество, ед."
        chart.y_axis.numFmt = "#,##0"
        chart.x_axis.title = "Период"

        data = Reference(
            worksheet,
            min_col=value_column,
            min_row=header_row,
            max_row=last_data_row,
        )
        categories = Reference(
            worksheet,
            min_col=category_column,
            min_row=first_data_row,
            max_row=last_data_row,
        )
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)

        self._style_chart_series(chart, line=True)
        return chart

    def _style_chart_series(self, chart: object, line: bool = False) -> None:
        if not chart.series:
            return
        series = chart.series[0]
        series.graphicalProperties.solidFill = self.theme.accent
        series.graphicalProperties.line.solidFill = self.theme.accent
        if line:
            series.graphicalProperties.line.width = 26000
            series.marker.symbol = "circle"
            series.marker.size = 6
            series.marker.graphicalProperties.solidFill = self.theme.white
            series.marker.graphicalProperties.line.solidFill = self.theme.accent

    def _style_block(
        self,
        worksheet: object,
        min_row: int,
        max_row: int,
        min_col: int,
        max_col: int,
        fill: str,
    ) -> None:
        pattern = PatternFill("solid", fgColor=fill)
        for row in worksheet.iter_rows(
            min_row=min_row,
            max_row=max_row,
            min_col=min_col,
            max_col=max_col,
        ):
            for cell in row:
                cell.fill = pattern

    def _finish_layout(self, worksheet: object, last_row: int) -> None:
        worksheet.print_area = f"A1:N{last_row}"
        worksheet.print_title_rows = "1:9"
        worksheet.page_setup.orientation = "landscape"
        worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A4
        worksheet.page_setup.fitToWidth = 1
        worksheet.page_setup.fitToHeight = 0
        worksheet.sheet_properties.pageSetUpPr.fitToPage = True
        worksheet.page_margins.left = 0.25
        worksheet.page_margins.right = 0.25
        worksheet.page_margins.top = 0.4
        worksheet.page_margins.bottom = 0.4
        worksheet.oddFooter.center.text = "Страница &P из &N"
        worksheet.oddFooter.center.size = 9
        worksheet.oddFooter.center.color = self.theme.muted


def _as_decimal(value: object) -> Decimal:
    if isinstance(value, Decimal):
        return value
    return Decimal(str(value))


def _excel_number(value: Decimal) -> int | float:
    integral = value.to_integral_value()
    return int(integral) if value == integral else float(value)