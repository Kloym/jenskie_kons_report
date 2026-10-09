"""Сводный отчёт ЖК по месяцам. Python 3.10+.

Запуск окна: python zhk_monthly_report.py
Зависимости: python -m pip install openpyxl xlrd
Встраивание: from zhk_monthly_report import launch_monthly
              launch_monthly(parent)

Исходные книги открываются только для чтения. Месяц берётся из даты.
Колонка «Тариф» выгрузки уже содержит сумму за всё количество в строке.
Цена одной услуги = Тариф / Кол-во. Пустая или нулевая сумма при Кол-во > 0 восстанавливается как
единственная цена (год, месяц, код) * Кол-во с округлением HALF_UP.
Положительные суммы не умножаются повторно. Нулевые суммы не служат донорами.
Промежуточные строки хранятся во временной SQLite, не в списке в памяти.
"""
from __future__ import annotations

import hashlib
import importlib.util
import math
import os
import re
import sqlite3
import tempfile
from collections import Counter, defaultdict
from contextlib import contextmanager
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path


UNITS = (
    "ЦЖЗ на Лобненской",
    "ЦЖЗ на Петрозаводской",
    "ЦЖЗ на Планетной",
    "ЖК №10",
)
MONTHS = (
    "", "Январь", "Февраль", "Март", "Апрель", "Май", "Июнь",
    "Июль", "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь",
)
CENT = Decimal("0.01")
ZERO = Decimal("0")
SUPPORTED = {".xls", ".xlsx", ".xlsm"}
ALIASES = {
    "date": {"дата", "дата услуги", "дата оказания услуги", "дата выполнения"},
    "code": {"код услуги", "код мед услуги", "код медуслуги"},
    "quantity": {"кол во", "количество", "количество услуг"},
    "tariff": {"тариф", "тариф руб", "цена", "цена руб", "цена тариф руб"},
    "department": {"подразделение", "наименование подразделения", "отделение", "жк"},
    "name": {"наименование услуги", "название услуги", "услуга"},
    "case_id": {"id пумп", "ид пумп"},
}
REQUIRED = ("date", "code", "quantity", "tariff", "department")


class ReportError(Exception):
    """Ошибка входных данных, при которой итог нельзя считать полным."""


class Cancelled(ReportError):
    pass


def ensure_dependencies(require_xls=False):
    missing = [name for name in (["openpyxl", "xlrd"] if require_xls else ["openpyxl"])
               if importlib.util.find_spec(name) is None]
    if missing:
        raise ReportError(
            "Не установлены библиотеки: " + ", ".join(missing)
            + ". Выполните: python -m pip install openpyxl xlrd"
        )


def clean_text(value):
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def normalize_key(value):
    return re.sub(r"[^0-9a-zа-я]+", " ", clean_text(value).casefold().replace("ё", "е")).strip()


def identify_department(value):
    key = normalize_key(value)
    women = "женск" in key or re.search(r"(?:^| )(?:жк|цжз)(?=$| |[0-9])", key)
    if not women:
        return None
    for fragment, unit in zip(("лобненск", "петрозаводск", "планетн"), UNITS):
        if fragment in key:
            return unit
    match = re.search(r"(?:женская консультация|жк)\s*([0-9]+)\s*([нh])?(?![0-9a-zа-я])", key)
    if match:
        number = int(match[1])
        suffix = match[2]
        if number == 13:
            return UNITS[1]
        if suffix:
            return None
        if number in (3, 5, 6, 7):
            return UNITS[2]
        if number == 10:
            return UNITS[3]
    return None


def decimal_number(value, label):
    if isinstance(value, bool):
        raise ValueError(f"{label}: логическое значение вместо числа")
    text = clean_text(value).replace(" ", "").replace(",", ".")
    try:
        result = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"{label}: пустое или некорректное число") from exc
    if not result.is_finite():
        raise ValueError(f"{label}: NaN/Infinity недопустимы")
    return result


def service_code(value):
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, (int, float, Decimal)):
        n = decimal_number(value, "Код услуги")
        if n < 0 or n != n.to_integral_value():
            raise ValueError("Код услуги должен быть целым числом или текстовым кодом")
        text = str(int(n))
    else:
        text = clean_text(value).lstrip("'")
    return text.zfill(6) if text.isdecimal() else text.upper()


def _case_key_text(value):
    text = clean_text(value)
    # Один ID может быть сохранён числом в одной выгрузке и текстом в другой.
    if re.fullmatch(r"[0-9]+(?:\.[0-9]+)?(?:[eE][+-]?[0-9]+)?", text):
        number = Decimal(text)
        if number.is_finite() and number == number.to_integral_value():
            return str(int(number))
    return text


def parse_date(value, epoch):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        from openpyxl.utils.datetime import from_excel
        try:
            parsed = from_excel(value, epoch=epoch)
            if isinstance(parsed, datetime):
                return parsed.date()
            if isinstance(parsed, date):
                return parsed
        except (ValueError, TypeError, OverflowError):
            pass
        raise ValueError("Некорректная числовая дата")
    text = clean_text(value)
    for fmt in ("%d.%m.%Y", "%d.%m.%y", "%Y-%m-%d", "%d/%m/%Y",
                "%d-%m-%Y", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError("Не удалось распознать дату услуги")


@dataclass
class Totals:
    quantity: Decimal = ZERO
    money: Decimal = ZERO
    rows: int = 0
    restored_rows: int = 0
    restored_money: Decimal = ZERO


@dataclass
class SourceStats:
    path: Path
    sheet: str
    rows: int = 0
    included: int = 0
    excluded: int = 0
    outside: int = 0
    empty_services: int = 0
    months: set = field(default_factory=set)
    min_date: date | None = None
    max_date: date | None = None


@dataclass
class Restoration:
    file: str
    sheet: str
    row: int
    service_date: date
    unit: str
    code: str
    tariff: Decimal | None  # Цена одной услуги, не исходная сумма строки.
    quantity: Decimal
    donor: str
    amount: Decimal
    original_amount: Decimal | None = None


@dataclass
class Report:
    year: int
    start_month: int
    end_month: int
    groups: dict = field(default_factory=lambda: defaultdict(Totals))
    loaded_months: set = field(default_factory=set)
    month_dates: dict = field(default_factory=lambda: defaultdict(set))
    sources: list = field(default_factory=list)
    restorations: list = field(default_factory=list)
    warnings: list = field(default_factory=list)
    excluded: Counter = field(default_factory=Counter)

    @property
    def months(self):
        return list(range(self.start_month, self.end_month + 1))


class _XlsSheet:
    def __init__(self, sheet, datemode):
        self.sheet = sheet
        self.title = sheet.name
        self.max_row = sheet.nrows
        self.datemode = datemode

    def iter_values(self, min_row=1, max_row=None):
        import xlrd
        stop = min(max_row or self.sheet.nrows, self.sheet.nrows)
        for i in range(min_row - 1, stop):
            values = []
            for cell in self.sheet.row(i):
                if cell.ctype == xlrd.XL_CELL_DATE:
                    values.append(xlrd.xldate_as_datetime(cell.value, self.datemode))
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    values.append(bool(cell.value))
                elif cell.ctype == xlrd.XL_CELL_ERROR:
                    values.append("#ОШИБКА_EXCEL")
                else:
                    values.append(cell.value)
            yield values


class _XlsxSheet:
    def __init__(self, sheet):
        self.sheet = sheet
        self.title = sheet.title
        self.max_row = sheet.max_row

    def iter_values(self, min_row=1, max_row=None):
        return self.sheet.iter_rows(min_row=min_row, max_row=max_row, values_only=True)


@contextmanager
def _open_source(path):
    if path.suffix.lower() == ".xls":
        import xlrd
        book = xlrd.open_workbook(str(path), on_demand=True)
        try:
            epoch = datetime(1904, 1, 1) if book.datemode else datetime(1899, 12, 30)
            yield [_XlsSheet(book.sheet_by_index(i), book.datemode) for i in range(book.nsheets)], epoch
        finally:
            book.release_resources()
    else:
        from openpyxl import load_workbook
        # Явно закрываем дескриптор и при ошибке/отмене посреди итерации.
        with path.open("rb") as stream:
            book = load_workbook(stream, read_only=True, data_only=True, keep_links=False)
            try:
                yield [_XlsxSheet(sheet) for sheet in book.worksheets], book.epoch
            finally:
                book.close()


def _find_header(sheet):
    # Завершаем короткий итератор, прежде чем запускать основной поток строк.
    header_rows = list(sheet.iter_values(max_row=30))
    for row_no, row in enumerate(header_rows, 1):
        keys = [normalize_key(v) for v in row]
        columns = {}
        for field_name, aliases in ALIASES.items():
            matches = [i for i, value in enumerate(keys) if value in aliases]
            if len(matches) == 1:
                columns[field_name] = matches[0]
        if all(name in columns for name in REQUIRED):
            return row_no, columns
    return None


def _cancelled(cancel):
    if cancel is not None and cancel.is_set():
        raise Cancelled("Обработка отменена. Исходные файлы не изменены.")


def _file_hash(path, cancel):
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            _cancelled(cancel)
            digest.update(block)
    return digest.digest()


def build_report(paths, year, start_month, end_month, progress=None, cancel=None):
    """Читает файлы, проверяет данные, восстанавливает тарифы и суммирует услуги."""
    if not (1900 <= year <= 9999 and 1 <= start_month <= end_month <= 12):
        raise ReportError("Укажите год и месяцы: начало не позже окончания.")
    paths = [Path(p).expanduser().resolve() for p in paths]
    if not paths:
        raise ReportError("Добавьте хотя бы один файл выгрузки.")
    ensure_dependencies(any(p.suffix.lower() == ".xls" for p in paths))
    report = Report(year, start_month, end_month)
    fingerprints, unique_paths = {}, []
    tell = progress or (lambda message: None)
    for path in paths:
        _cancelled(cancel)
        if not path.is_file() or path.suffix.lower() not in SUPPORTED:
            raise ReportError(f"Не найден поддерживаемый Excel-файл: {path}")
        fingerprint = _file_hash(path, cancel)
        if fingerprint in fingerprints:
            report.warnings.append(
                f"Повторный файл {path.name} не учтён: содержимое совпадает с {fingerprints[fingerprint].name}."
            )
            continue
        fingerprints[fingerprint] = path
        unique_paths.append(path)

    # Только минимальные поля услуг, без ФИО пациентов и полисов.
    tariffs = defaultdict(dict)
    with tempfile.TemporaryDirectory(prefix="zhk_monthly_") as temp:
        connection = sqlite3.connect(str(Path(temp) / "rows.sqlite"))
        try:
            connection.execute("CREATE TABLE rows (source INTEGER, row_no INTEGER, day TEXT, month INTEGER, unit TEXT, code TEXT, quantity TEXT, tariff TEXT)")
            connection.execute("CREATE TABLE cases (key BLOB PRIMARY KEY, source INTEGER) WITHOUT ROWID")
            for file_index, path in enumerate(unique_paths, 1):
                tell(f"Чтение {file_index}/{len(unique_paths)}: {path.name}")
                stat_before = (path.stat().st_size, path.stat().st_mtime_ns)
                recognized = 0
                try:
                    with _open_source(path) as (sheets, epoch):
                        for sheet in sheets:
                            _cancelled(cancel)
                            found = _find_header(sheet)
                            if found is None:
                                report.warnings.append(f"{path.name}, лист {sheet.title}: таблица услуг не найдена; лист пропущен.")
                                continue
                            recognized += 1
                            header_row, columns = found
                            source = SourceStats(path, sheet.title)
                            source_id = len(report.sources)
                            report.sources.append(source)
                            if path.suffix.lower() == ".xls" and sheet.max_row >= 65536:
                                report.warnings.append(f"{path.name}, лист {sheet.title}: достигнут предел 65 536 строк XLS. Проверьте полноту выгрузки.")
                            if "case_id" not in columns:
                                report.warnings.append(f"{path.name}, лист {sheet.title}: нет ID ПУМП; пересечения обращений между разными файлами не проверяются.")
                            for row_no, row in enumerate(sheet.iter_values(min_row=header_row + 1), header_row + 1):
                                if row_no % 3000 == 0:
                                    _cancelled(cancel)
                                    tell(f"{path.name}: прочитано {row_no - header_row:,} строк".replace(",", " "))
                                if not any(clean_text(v) for v in row):
                                    continue
                                source.rows += 1
                                def value(key):
                                    i = columns.get(key)
                                    return row[i] if i is not None and i < len(row) else None
                                try:
                                    dep = clean_text(value("department"))
                                    if not dep:
                                        raise ValueError("Не заполнено подразделение")
                                    unit = identify_department(dep)
                                    if unit is None:
                                        source.excluded += 1
                                        key = normalize_key(dep)
                                        if "женск" in key or re.search(r"(?:^| )(?:жк|цжз)(?=$| |[0-9])", key):
                                            dt = parse_date(value("date"), epoch)
                                            if dt.year == year and start_month <= dt.month <= end_month:
                                                report.excluded[dep] += 1
                                        continue
                                    dt = parse_date(value("date"), epoch)
                                    if dt.year != year or not start_month <= dt.month <= end_month:
                                        source.outside += 1
                                        continue
                                    if not any(clean_text(value(k)) for k in ("code", "quantity", "tariff", "name")):
                                        source.empty_services += 1
                                        continue
                                    code = service_code(value("code"))
                                    if not code:
                                        raise ValueError("Не заполнен код услуги")
                                    quantity = decimal_number(value("quantity"), "Кол-во")
                                    if quantity < 0 or quantity != quantity.to_integral_value():
                                        raise ValueError("Количество должно быть целым числом >= 0")
                                    raw_tariff = value("tariff")
                                    tariff = decimal_number(raw_tariff, "Тариф") if clean_text(raw_tariff) else None
                                    if tariff is not None and tariff < 0:
                                        raise ValueError("Тариф должен быть >= 0")
                                    case_id = _case_key_text(value("case_id"))
                                    if case_id:
                                        case_key = hashlib.sha256((dt.isoformat() + "|" + case_id).encode()).digest()
                                        previous = connection.execute("SELECT source FROM cases WHERE key=?", (case_key,)).fetchone()
                                        if previous is not None and previous[0] != source_id:
                                            other = report.sources[previous[0]]
                                            raise ValueError(
                                                f"Обращение уже встречалось в {other.path.name}, лист {other.sheet}. "
                                                "Выгрузки пересекаются: загрузите непересекающиеся файлы, чтобы не удвоить суммы."
                                            )
                                        connection.execute("INSERT OR IGNORE INTO cases VALUES (?,?)", (case_key, source_id))
                                    connection.execute("INSERT INTO rows VALUES (?,?,?,?,?,?,?,?)", (
                                        source_id, row_no, dt.isoformat(), dt.month, unit, code,
                                        str(quantity), None if tariff is None else str(tariff),
                                    ))
                                    if tariff is not None and tariff > 0 and quantity > 0:
                                        # В выгрузке «Тариф» — сумма за всё количество.
                                        # Сравниваем цены за единицу без раннего округления.
                                        unit_price = tariff / quantity
                                        donor = (f"{path.name}, {sheet.title}, строка {row_no}; "
                                                 f"{tariff} / {quantity} = {unit_price} руб./услугу")
                                        tariffs[(dt.month, code)].setdefault(unit_price, donor)
                                    source.included += 1
                                    source.months.add(dt.month)
                                    source.min_date = min(source.min_date, dt) if source.min_date else dt
                                    source.max_date = max(source.max_date, dt) if source.max_date else dt
                                    report.month_dates[dt.month].add(dt)
                                except (ValueError, TypeError, IndexError, InvalidOperation) as exc:
                                    raise ReportError(f"{path.name}, лист {sheet.title}, строка {row_no}: {exc}") from exc
                    if not recognized:
                        raise ReportError(f"{path.name}: не найдена таблица с колонками Дата, Код услуги, Кол-во, Тариф, Подразделение.")
                except (ReportError, Cancelled):
                    raise
                except Exception as exc:
                    raise ReportError(f"Не удалось прочитать {path.name}: {exc}") from exc
                if stat_before != (path.stat().st_size, path.stat().st_mtime_ns):
                    raise ReportError(f"{path.name} изменён во время чтения. Повторите расчёт.")
                connection.commit()

            tell("Восстановление тарифов и расчёт итогов…")
            unresolved = []
            unresolved_count = 0
            for i, record in enumerate(connection.execute("SELECT source,row_no,day,month,unit,code,quantity,tariff FROM rows"), 1):
                if i % 3000 == 0:
                    _cancelled(cancel)
                source_id, row_no, day, month, unit, code, quantity, raw_tariff = record
                quantity = Decimal(quantity)
                original_amount = None if raw_tariff is None else Decimal(raw_tariff)
                restored = original_amount is None or (original_amount == 0 and quantity > 0)
                source = report.sources[source_id]
                if restored and quantity == 0:
                    tariff = None
                    amount = ZERO
                    donor = "Количество равно 0; сумма строки = 0. Цена одной услуги не определяется."
                elif restored:
                    options = tariffs.get((month, code), {})
                    if len(options) != 1:
                        unresolved_count += 1
                        if len(unresolved) < 12:
                            reason = ("нет ненулевого тарифа при количестве > 0" if not options else
                                      "разные тарифы за 1 услугу (Тариф / Кол-во): " + ", ".join(str(v) for v in sorted(options)))
                            unresolved.append(f"{source.path.name}, {source.sheet}, строка {row_no}, код {code}, {MONTHS[month]}: {reason}")
                        continue
                    tariff, donor = next(iter(options.items()))
                    amount = (quantity * tariff).quantize(CENT, rounding=ROUND_HALF_UP)
                else:
                    amount = original_amount.quantize(CENT, rounding=ROUND_HALF_UP)
                if restored:
                    report.restorations.append(Restoration(source.path.name, source.sheet, row_no, date.fromisoformat(day), unit, code, tariff, quantity, donor, amount, original_amount))
                total = report.groups[(month, unit)]
                total.quantity += quantity
                total.money += amount
                total.rows += 1
                if restored:
                    total.restored_rows += 1
                    total.restored_money += amount
                report.loaded_months.add(month)
            if unresolved_count:
                extra = f"\nПоказаны первые {len(unresolved)}." if unresolved_count > len(unresolved) else ""
                raise ReportError(
                    f"Не удалось однозначно восстановить {unresolved_count} пустых или нулевых тарифов. "
                    "Итог не сохранён, чтобы не занизить сумму.\n" + "\n".join(unresolved) + extra
                )
        finally:
            connection.close()
    _cancelled(cancel)
    if not report.loaded_months:
        raise ReportError(f"Нет услуг нужных ЖК за {year} год, {MONTHS[start_month]}–{MONTHS[end_month]}. Месяц определяется по дате, а не по имени файла.")
    missing = [MONTHS[m] for m in report.months if m not in report.loaded_months]
    if missing:
        report.warnings.append("В выбранных выгрузках нет услуг за месяцы: " + ", ".join(missing) + ". В отчёте они отмечены как «Нет данных».")
    if report.excluded:
        report.warnings.append("Не включены нераспознанные ЖК: " + "; ".join(f"{name}: {count} строк" for name, count in report.excluded.items()))
    if any(s.empty_services for s in report.sources):
        report.warnings.append(f"Пропущено строк без услуги: {sum(s.empty_services for s in report.sources)}. В них одновременно пусты код, название, количество и тариф.")
    count = sum(t.rows for t in report.groups.values())
    money = sum((t.money for t in report.groups.values()), ZERO)
    quantity = sum((t.quantity for t in report.groups.values()), ZERO)
    tell(f"Учтено {count:,} строк; услуг: {quantity:,}; сумма: {money:,.2f} руб. Восстановлено тарифов: {len(report.restorations)}.".replace(",", " "))
    return report


def save_report(report, output, source_paths, overwrite=False):
    """Атомарное сохранение с защитой исходных файлов и существующего результата."""
    output = Path(output).expanduser().resolve()
    if output.suffix.lower() != ".xlsx":
        raise ReportError("Сохраните результат с расширением .xlsx.")
    for source in source_paths:
        source = Path(source).expanduser().resolve()
        if output == source or (output.exists() and source.exists() and os.path.samefile(output, source)):
            raise ReportError("Нельзя перезаписывать исходную выгрузку итоговым отчётом.")
    if output.exists() and not overwrite:
        raise ReportError("Выходной файл уже существует. Выберите другое имя.")
    if not output.parent.is_dir():
        raise ReportError("Папка для сохранения не существует.")
    book = make_workbook(report)
    temp_name = None
    try:
        descriptor, temp_name = tempfile.mkstemp(prefix=".zhk_", suffix=".xlsx", dir=output.parent)
        os.close(descriptor)
        book.save(temp_name)
        if not overwrite and output.exists():
            raise ReportError("Файл с таким именем появился во время сохранения. Выберите другое имя.")
        os.replace(temp_name, output)
        temp_name = None
    except PermissionError as exc:
        raise ReportError("Не удалось сохранить файл. Закройте его в Excel и проверьте доступ к папке.") from exc
    finally:
        book.close()
        if temp_name:
            Path(temp_name).unlink(missing_ok=True)
    return output


"""Workbook presentation for the monthly report (lazy openpyxl imports)."""

from decimal import Decimal
from pathlib import Path

_XL_INK = "20334D"
_XL_MUTED = "65758A"
_XL_RULE = "D6DFE9"
_XL_LIGHT = "F2F5F9"
_XL_BLUE = "244769"
_XL_GROUP_COLORS = ("30688D", "39776F", "69618B", "917344")
_XL_GROUP_LIGHT = ("EDF4FA", "EDF6F3", "F2F0F8", "F8F4EC")
_XL_MONEY = '#,##0.00;[Red](#,##0.00);0.00'
_XL_COUNT = '#,##0;[Red](#,##0);0'
_XL_PCT = '+0.0%;[Red]-0.0%;0.0%'


def _xl_num(value):
    """Convert only at the Excel boundary; aggregation has already used Decimal."""
    value = Decimal(value)
    return int(value) if value == value.to_integral_value() else float(value)


def _xl_cell(ws, row, col, value, number_format=None):
    from openpyxl.styles import Alignment, Font
    cell = ws.cell(row, col)
    if isinstance(value, str):
        cell.value = value[:32767]
        cell.data_type = "s"
    else:
        cell.value = value
    cell.font = Font(name="Arial", size=10, color=_XL_INK)
    cell.alignment = Alignment(vertical="center", horizontal="left" if isinstance(value, str) else "right")
    if number_format:
        cell.number_format = number_format
    return cell


def _xl_base(ws, widths, title, last_col):
    from openpyxl.styles import Border, Font, Side
    from openpyxl.utils import get_column_letter
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_properties.tabColor = _XL_BLUE
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 8
    ws.row_dimensions[2].height = 29
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    c = _xl_cell(ws, 2, 1, title)
    c.font = Font(name="Arial", size=16, bold=True, color=_XL_INK)
    for col in range(1, last_col + 1):
        ws.cell(3, col).border = Border(bottom=Side(style="thin", color=_XL_RULE))
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = ws.page_margins.right = 0.3
    ws.page_margins.top = ws.page_margins.bottom = 0.35
    ws.oddFooter.left.text = "Сводный отчёт ЖК"
    ws.oddFooter.right.text = "Страница &P из &N"
    ws.oddFooter.left.size = ws.oddFooter.right.size = 8


def _xl_note(ws, row, text, last_col, height=30, warning=False):
    from openpyxl.styles import Alignment, Font, PatternFill
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    c = _xl_cell(ws, row, 1, text)
    c.font = Font(name="Arial", size=10, color="8A5B16" if warning else _XL_MUTED)
    c.alignment = Alignment(vertical="center", wrap_text=True)
    if warning:
        c.fill = PatternFill("solid", fgColor="FFF3D9")
    ws.row_dimensions[row].height = height


def _xl_header(ws, row, names, start_col=1, color=_XL_BLUE, height=34):
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    for col, name in enumerate(names, start_col):
        c = _xl_cell(ws, row, col, name)
        c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=color)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(right=Side(style="thin", color="FFFFFF"))
    ws.row_dimensions[row].height = height


def _xl_total_style(ws, row, first_col, last_col):
    from openpyxl.styles import Font, PatternFill, Border, Side
    for col in range(first_col, last_col + 1):
        c = ws.cell(row, col)
        c.font = Font(name="Arial", size=10, bold=True, color=_XL_INK)
        c.fill = PatternFill("solid", fgColor="E3EAF2")
        c.border = Border(top=Side(style="thin", color="A7B7CA"))
    ws.row_dimensions[row].height = 27


def _xl_chart(ws, data_col, first_row, last_row, anchor, title, unit, color, line=False):
    """Native chart with explicit caches, so previews also have plotted values."""
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.chart.data_source import AxDataSource, NumData, NumVal, StrData, StrRef, StrVal
    chart = LineChart() if line else BarChart()
    chart.title = title
    chart.y_axis.title = unit
    chart.legend = None
    chart.style = 13
    chart.width = 15.8
    chart.height = 8.8
    chart.display_blanks = "gap"
    chart.y_axis.numFmt = "#,##0"
    chart.add_data(Reference(ws, min_col=data_col, min_row=first_row - 1, max_row=last_row), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=first_row, max_row=last_row))
    count = last_row - first_row + 1
    series = chart.series[0]
    values = [NumVal(idx=i, v=ws.cell(row, data_col).value)
              for i, row in enumerate(range(first_row, last_row + 1))
              if isinstance(ws.cell(row, data_col).value, (int, float))]
    series.val.numRef.numCache = NumData(formatCode="General", ptCount=count, pt=values)
    labels = [StrVal(idx=i, v=str(ws.cell(row, 1).value))
              for i, row in enumerate(range(first_row, last_row + 1))]
    quoted_sheet = "'" + ws.title.replace("'", "''") + "'"
    series.cat = AxDataSource(strRef=StrRef(
        f=f"{quoted_sheet}!A{first_row}:A{last_row}",
        strCache=StrData(ptCount=count, pt=labels)))
    series.graphicalProperties.line.solidFill = color
    if line:
        series.graphicalProperties.line.width = 26000
        series.marker.symbol = "circle"
        series.marker.size = 5
        series.marker.graphicalProperties.solidFill = color
        series.marker.graphicalProperties.line.solidFill = color
        series.smooth = False
    else:
        series.graphicalProperties.solidFill = color
        chart.gapWidth = 65
    ws.add_chart(chart, anchor)


def make_workbook(report):
    """Return a styled snapshot. Input files and patient data are not copied."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.workbook.properties import CalcProperties
    book = Workbook()
    book.properties.title = f"Сводный отчёт ЖК за {report.year} год"
    book.properties.subject = "Количество и стоимость услуг по месяцам и подразделениям"
    book.properties.creator = "Сводный отчёт ЖК"
    book.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True)
    months = list(report.months)
    zero = Decimal("0")
    month_quantity = {m: sum((report.groups[(m, u)].quantity for u in UNITS if (m, u) in report.groups), zero) for m in months}
    month_money = {m: sum((report.groups[(m, u)].money for u in UNITS if (m, u) in report.groups), zero) for m in months}
    quantity_total = sum(month_quantity.values(), zero)
    money_total = sum(month_money.values(), zero)
    restored_rows = sum(t.restored_rows for (m, _), t in report.groups.items() if m in months)
    restored_money = sum((t.restored_money for (m, _), t in report.groups.items() if m in months), zero)
    period = f"{MONTHS[report.start_month]} — {MONTHS[report.end_month]}, {report.year} год"

    ws = book.active
    ws.title = "Итоги"
    _xl_base(ws, (16, 12, 19, 12, 19, 12, 19, 12, 19, 13, 21, 25), "Услуги женских консультаций", 12)
    _xl_note(ws, 4, period + ". Количество — сумма «Кол-во»; стоимость — сумма колонки «Тариф» (уже за всё количество), руб.", 12, 25)
    _xl_header(ws, 6, ("Месяц", "", "", "", "", "", "", "", "", "Всего", "", "Данные месяца"), height=28)
    _xl_header(ws, 7, ("", "Услуги, ед.", "Сумма, руб.", "Услуги, ед.", "Сумма, руб.", "Услуги, ед.", "Сумма, руб.", "Услуги, ед.", "Сумма, руб.", "Услуги, ед.", "Сумма, руб.", ""), height=27)
    for col in (1, 12):
        ws.merge_cells(start_row=6, start_column=col, end_row=7, end_column=col)
    ws.merge_cells("J6:K6")
    for i, unit in enumerate(UNITS):
        first = 2 + 2 * i
        ws.merge_cells(start_row=6, start_column=first, end_row=6, end_column=first + 1)
        c = _xl_cell(ws, 6, first, unit)
        c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        c.alignment = Alignment(horizontal="center", vertical="center")
        for row in (6, 7):
            for col in (first, first + 1):
                ws.cell(row, col).fill = PatternFill("solid", fgColor=_XL_GROUP_COLORS[i])
    for row, month in enumerate(months, 8):
        present = month in report.loaded_months
        _xl_cell(ws, row, 1, MONTHS[month])
        ws.row_dimensions[row].height = 34
        for i, unit in enumerate(UNITS):
            totals = report.groups.get((month, unit))
            for col, value, fmt in ((2 + i * 2, totals.quantity if totals else zero, _XL_COUNT),
                                    (3 + i * 2, totals.money if totals else zero, _XL_MONEY)):
                c = _xl_cell(ws, row, col, _xl_num(value) if present else None, fmt)
                c.fill = PatternFill("solid", fgColor=_XL_GROUP_LIGHT[i] if present else "F4F5F7")
        _xl_cell(ws, row, 10, _xl_num(month_quantity[month]) if present else None, _XL_COUNT)
        _xl_cell(ws, row, 11, _xl_num(month_money[month]) if present else None, _XL_MONEY)
        dates = sorted(getattr(report, "month_dates", {}).get(month, ()))
        if dates:
            status = f"{dates[0]:%d.%m}–{dates[-1]:%d.%m}\nДней с услугами: {len(dates)}"
        else:
            status = "Есть строки" if present else "Нет данных"
        c = _xl_cell(ws, row, 12, status)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.font = Font(name="Arial", size=10, color=_XL_MUTED if present else "996015")
        if not present:
            c.fill = PatternFill("solid", fgColor="FFF3D9")
        for col in (10, 11):
            ws.cell(row, col).font = Font(name="Arial", size=10, bold=True, color=_XL_INK)
    last_row = 7 + len(months)
    total_row = last_row + 1
    _xl_cell(ws, total_row, 1, "Итого")
    for i, unit in enumerate(UNITS):
        totals = [report.groups[(m, unit)] for m in months if (m, unit) in report.groups]
        _xl_cell(ws, total_row, 2 + 2 * i, _xl_num(sum((t.quantity for t in totals), zero)), _XL_COUNT)
        _xl_cell(ws, total_row, 3 + 2 * i, _xl_num(sum((t.money for t in totals), zero)), _XL_MONEY)
    _xl_cell(ws, total_row, 10, _xl_num(quantity_total), _XL_COUNT)
    _xl_cell(ws, total_row, 11, _xl_num(money_total), _XL_MONEY)
    _xl_cell(ws, total_row, 12, f"Месяцев с данными: {len(set(months) & report.loaded_months)}")
    _xl_total_style(ws, total_row, 1, 12)
    _xl_note(ws, total_row + 2, "«Нет данных» — в выбранных файлах нет учтённых услуг за месяц. Ноль у отдельного подразделения означает отсутствие его строк в загруженных данных. Итоги включают только загруженные строки.", 12, 32)
    _xl_note(ws, total_row + 3, "Даты и число дней с услугами показывают фактическое покрытие выгрузки и не подтверждают её полноту за месяц. При разных периодах покрытия динамику следует сравнивать с учётом этого ограничения.", 12, 32)
    _xl_note(ws, total_row + 4, "Объединение с января: ЖК №3, 5, 6, 7 → ЦЖЗ на Планетной; ЖК №13 и №13Н → ЦЖЗ на Петрозаводской. Месяц определяется по дате услуги, а не по имени файла.", 12, 32)
    amount_text = f"{restored_money:,.2f}".replace(",", " ").replace(".", ",")
    _xl_note(ws, total_row + 5, f"Восстановлено пустых и нулевых тарифов: {restored_rows}. Стоимость строк с подстановкой: {amount_text} руб. Она уже включена в общую сумму." + (" Подробности — на листе «Тарифы»." if report.restorations else ""), 12, 30, bool(restored_rows))
    ws.freeze_panes = "B8"
    ws.print_title_rows = "1:7"
    ws.print_area = f"A1:L{total_row + 5}"
    ws.page_setup.fitToHeight = 1

    analytics = book.create_sheet("Аналитика")
    _xl_base(analytics, (16, 14, 20, 18, 17, 17, 3, 30, 14, 20, 13), "Динамика услуг и структура суммы", 11)
    _xl_note(analytics, 4, period + ". Показатели рассчитаны по учтённым строкам.", 11, 25)
    _xl_header(analytics, 6, ("Месяц", "Услуги, ед.", "Сумма, руб.", "На услугу, руб.", "Услуги к пред. мес.", "Сумма к пред. мес."))
    for row, month in enumerate(months, 7):
        present = month in report.loaded_months
        _xl_cell(analytics, row, 1, MONTHS[month])
        q, money = month_quantity[month], month_money[month]
        _xl_cell(analytics, row, 2, _xl_num(q) if present else None, _XL_COUNT)
        _xl_cell(analytics, row, 3, _xl_num(money) if present else None, _XL_MONEY)
        _xl_cell(analytics, row, 4, _xl_num(money / q) if present and q else None, _XL_MONEY)
        previous = month - 1
        adjacent = present and previous in report.loaded_months and previous in month_quantity
        for col, current, values in ((5, q, month_quantity), (6, money, month_money)):
            old = values.get(previous, zero)
            growth = _xl_num(current / old - 1) if adjacent and old else None
            _xl_cell(analytics, row, col, growth, _XL_PCT)
        analytics.row_dimensions[row].height = 25
        for col in range(1, 7):
            if not present or row % 2:
                analytics.cell(row, col).fill = PatternFill("solid", fgColor="F4F5F7" if not present else "F3F6FA")
    analytics_total = 7 + len(months)
    for col, value, fmt in ((1, "Итого", None), (2, _xl_num(quantity_total), _XL_COUNT),
                            (3, _xl_num(money_total), _XL_MONEY),
                            (4, _xl_num(money_total / quantity_total) if quantity_total else None, _XL_MONEY)):
        _xl_cell(analytics, analytics_total, col, value, fmt)
    _xl_total_style(analytics, analytics_total, 1, 6)
    _xl_header(analytics, 6, ("Подразделение", "Услуги, ед.", "Сумма, руб.", "Доля суммы"), 8)
    for row, unit in enumerate(UNITS, 7):
        totals = [report.groups[(m, unit)] for m in months if (m, unit) in report.groups]
        q = sum((t.quantity for t in totals), zero)
        money = sum((t.money for t in totals), zero)
        for col, value, fmt in ((8, unit, None), (9, _xl_num(q), _XL_COUNT), (10, _xl_num(money), _XL_MONEY),
                                (11, _xl_num(money / money_total) if money_total else None, "0.0%")):
            c = _xl_cell(analytics, row, col, value, fmt)
            c.fill = PatternFill("solid", fgColor=_XL_GROUP_LIGHT[row - 7])
    for col, value, fmt in ((8, "Итого", None), (9, _xl_num(quantity_total), _XL_COUNT),
                            (10, _xl_num(money_total), _XL_MONEY), (11, 1 if money_total else None, "0.0%")):
        _xl_cell(analytics, 11, col, value, fmt)
    _xl_total_style(analytics, 11, 8, 11)
    note_row = max(analytics_total + 2, 14)
    _xl_note(analytics, note_row, "Изменение к предыдущему месяцу показано только при наличии двух соседних месяцев и ненулевой базы. Пустая ячейка означает, что показатель не рассчитан. «На услугу» = общая сумма ÷ количество услуг.", 11, 32)
    _xl_note(analytics, note_row + 1, "Графики отражают объём выгрузок. Пропущенные месяцы не считаются нулевыми; наличие строк не означает полноту месяца. Все показатели включают строки с восстановленными тарифами.", 11, 32)
    chart_row = note_row + 3
    _xl_chart(analytics, 3, 7, analytics_total - 1, f"A{chart_row}", "Стоимость услуг по месяцам", "Рубли", _XL_BLUE)
    _xl_chart(analytics, 2, 7, analytics_total - 1, f"G{chart_row}", "Количество услуг по месяцам", "Услуги, ед.", "39776F", line=True)
    for column in ("E", "F"):
        analytics.conditional_formatting.add(f"{column}7:{column}{analytics_total - 1}",
            CellIsRule(operator="lessThan", formula=["0"], font=Font(name="Arial", color="9B3940")))
    analytics.print_area = f"A1:K{chart_row + 18}"
    analytics.page_setup.fitToHeight = 1

    sources = book.create_sheet("Источники")
    _xl_base(sources, (40, 23, 14, 14, 15, 17, 17, 30, 15, 15), "Источники и обработка данных", 10)
    sources.sheet_properties.tabColor = "8797AA"
    _xl_note(sources, 4, "Исходные книги не изменяются. В отчёт перенесены только сводные показатели и технические сведения; персональные данные пациентов не копируются.", 10, 30)
    _xl_header(sources, 6, ("Файл", "Лист", "Строк просмотрено", "Учтено строк", "Другие отделения", "Вне периода", "Без услуги", "Месяцы услуг", "Первая дата", "Последняя дата"), height=38)
    for row, source in enumerate(report.sources, 7):
        values = (Path(source.path).name, source.sheet, source.rows, source.included, source.excluded,
                  source.outside, source.empty_services, ", ".join(MONTHS[m] for m in sorted(source.months)),
                  source.min_date, source.max_date)
        for col, value in enumerate(values, 1):
            c = _xl_cell(sources, row, col, value, "dd.mm.yyyy" if col in (9, 10) else _XL_COUNT if col in range(3, 8) else None)
            if col in (1, 2, 8):
                c.alignment = Alignment(vertical="center", wrap_text=True)
            if row % 2:
                c.fill = PatternFill("solid", fgColor=_XL_LIGHT)
        sources.row_dimensions[row].height = 34
    source_end = 6 + len(report.sources)
    if report.sources:
        sources.auto_filter.ref = f"A6:J{source_end}"
    row = source_end + 2
    _xl_note(sources, row, "Строки, количество и сумма — разные показатели: одна строка может содержать несколько услуг. Даты ниже и на листе «Итоги» относятся к учтённым услугам.", 10, 30)
    row += 2
    _xl_header(sources, row, ("Файлы: полный путь",), height=25)
    for path in dict.fromkeys(str(s.path) for s in report.sources):
        row += 1
        _xl_note(sources, row, path, 10, 30)
    row += 2
    _xl_header(sources, row, ("Сообщения проверки",), height=25)
    for message in report.warnings or ["Дополнительных замечаний при чтении файлов нет."]:
        row += 1
        _xl_note(sources, row, str(message), 10, 38, bool(report.warnings))
    if report.excluded:
        row += 2
        _xl_header(sources, row, ("Исключённое подразделение", "Строк"))
        for department, count in sorted(report.excluded.items(), key=lambda item: (-item[1], item[0])):
            row += 1
            _xl_cell(sources, row, 1, department or "Без названия").alignment = Alignment(wrap_text=True, vertical="center")
            _xl_cell(sources, row, 2, count, _XL_COUNT)
            sources.row_dimensions[row].height = 35
    sources.freeze_panes = "C7"
    sources.print_title_rows = "1:6"
    sources.print_area = f"A1:J{row}"

    if report.restorations:
        tariffs = book.create_sheet("Тарифы")
        _xl_base(tariffs, (35, 22, 12, 15, 30, 16, 17, 15, 21, 50, 19), "Восстановленные пустые и нулевые тарифы", 11)
        _xl_note(tariffs, 4, "Пустые и нулевые суммы восстановлены по единственной положительной цене за тот же код, год и месяц. Цена = Тариф / Кол-во у источника; новая сумма = цена × количество. Пустая сумма при количестве 0 равна 0. Подстановки включены в итоги; исходники не изменены.", 11, 34)
        _xl_header(tariffs, 6, ("Файл", "Лист", "Строка Excel", "Дата услуги", "Подразделение", "Код услуги", "Цена за 1 услугу, руб.", "Услуги, ед.", "Восстановленная сумма, руб.", "Источник цены", "Исходное значение"), height=38)
        for row, item in enumerate(report.restorations, 7):
            values = (item.file, item.sheet, item.row, item.service_date, item.unit, item.code,
                      _xl_num(item.tariff) if item.tariff is not None else None,
                      _xl_num(item.quantity), _xl_num(item.amount), item.donor,
                      "Пусто" if item.original_amount is None else _xl_num(item.original_amount))
            for col, value in enumerate(values, 1):
                fmt = "dd.mm.yyyy" if col == 4 else "0.00######" if col == 7 else _XL_MONEY if col == 9 else _XL_COUNT if col in (3, 8) else None
                c = _xl_cell(tariffs, row, col, value, fmt)
                if col in (1, 2, 5, 10):
                    c.alignment = Alignment(vertical="center", wrap_text=True)
                if row % 2:
                    c.fill = PatternFill("solid", fgColor=_XL_LIGHT)
            tariffs.row_dimensions[row].height = 33
        end = 6 + len(report.restorations)
        tariffs.auto_filter.ref = f"A6:K{end}"
        tariffs.freeze_panes = "D7"
        tariffs.print_title_rows = "1:6"
        tariffs.print_area = f"A1:K{end}"
    book.active = 0
    return book



# -------------------- Окно и командная строка --------------------
import argparse
import queue
import subprocess
import sys
import threading


def _open_result_file(path):
    """Открывает только по явному нажатию пользователя."""
    if sys.platform == "win32":
        os.startfile(str(path))
    elif sys.platform == "darwin":
        subprocess.Popen(["open", str(path)])
    else:
        subprocess.Popen(["xdg-open", str(path)])


def launch_monthly(parent=None):
    """Окно отчёта. Без parent запускает приложение; с parent — дочернее окно."""
    try:
        import tkinter as tk
        from tkinter import filedialog, messagebox, ttk
    except ImportError as exc:
        raise ReportError(
            "Не установлен Tkinter. В Windows переустановите Python с компонентом "
            "Tcl/Tk; в Linux установите пакет python3-tk."
        ) from exc

    try:
        window = tk.Tk() if parent is None else tk.Toplevel(parent)
    except tk.TclError as exc:
        raise ReportError(
            "Не удалось открыть окно. Запустите программу в графической среде "
            "или используйте командную строку: --help."
        ) from exc
    window.title("ЖК · сводный отчёт по месяцам")
    width = max(700, min(1040, window.winfo_screenwidth() - 60))
    height = max(580, min(760, window.winfo_screenheight() - 80))
    window.geometry(f"{width}x{height}")
    window.minsize(min(780, width), min(620, height))
    window.configure(background="#eff4f8")
    window.columnconfigure(0, weight=1)
    window.rowconfigure(0, weight=1)

    style = ttk.Style(window)
    style.configure("Monthly.TFrame", background="#eff4f8")
    style.configure("Monthly.TLabel", background="#eff4f8", foreground="#19334a")
    style.configure("Monthly.Title.TLabel", font=("Arial", 21, "bold"), background="#eff4f8", foreground="#15354c")
    style.configure("Monthly.Muted.TLabel", background="#eff4f8", foreground="#546b7d")
    style.configure("Monthly.TLabelframe", background="#eff4f8")
    style.configure("Monthly.TLabelframe.Label", background="#eff4f8", foreground="#19334a", font=("Arial", 11, "bold"))
    style.configure("Monthly.Treeview", rowheight=27)
    style.configure("Monthly.Primary.TButton", font=("Arial", 10, "bold"), padding=(14, 10))

    root_frame = ttk.Frame(window, padding=(22, 17), style="Monthly.TFrame")
    root_frame.grid(sticky="nsew")
    root_frame.columnconfigure(0, weight=1)
    root_frame.rowconfigure(3, weight=2)
    root_frame.rowconfigure(6, weight=1)
    ttk.Label(root_frame, text="ЖК / Итоги по месяцам", style="Monthly.Title.TLabel").grid(sticky="w")
    ttk.Label(root_frame, text="Добавьте выгрузки, выберите период и получите одну сводную Excel-книгу.", style="Monthly.Muted.TLabel").grid(sticky="w", pady=(5, 12))

    period_frame = ttk.LabelFrame(root_frame, text="  1. Период отчёта  ", padding=12, style="Monthly.TLabelframe")
    period_frame.grid(row=2, column=0, sticky="ew", pady=(0, 12))
    now = datetime.now()
    year_var = tk.StringVar(master=window, value=str(now.year))
    start_var = tk.StringVar(master=window, value=MONTHS[1])
    end_var = tk.StringVar(master=window, value=MONTHS[now.month])
    ttk.Label(period_frame, text="Год", style="Monthly.TLabel").grid(row=0, column=0, sticky="w", padx=(0, 7))
    year_input = ttk.Spinbox(period_frame, from_=2000, to=2100, textvariable=year_var, width=7)
    year_input.grid(row=0, column=1, padx=(0, 22))
    ttk.Label(period_frame, text="С", style="Monthly.TLabel").grid(row=0, column=2, padx=(0, 7))
    from_input = ttk.Combobox(period_frame, textvariable=start_var, values=MONTHS[1:], state="readonly", width=13)
    from_input.grid(row=0, column=3, padx=(0, 17))
    ttk.Label(period_frame, text="По", style="Monthly.TLabel").grid(row=0, column=4, padx=(0, 7))
    to_input = ttk.Combobox(period_frame, textvariable=end_var, values=MONTHS[1:], state="readonly", width=13)
    to_input.grid(row=0, column=5)
    ttk.Label(period_frame, text="Месяц определяется по дате услуги внутри файла.", style="Monthly.Muted.TLabel").grid(row=1, column=0, columnspan=6, sticky="w", pady=(10, 0))

    files_frame = ttk.LabelFrame(root_frame, text="  2. Исходные файлы  ", padding=12, style="Monthly.TLabelframe")
    files_frame.grid(row=3, column=0, sticky="nsew", pady=(0, 9))
    files_frame.columnconfigure(0, weight=1)
    files_frame.rowconfigure(0, weight=1)
    files_tree = ttk.Treeview(files_frame, columns=("file", "folder"), show="headings", selectmode="extended", height=5, style="Monthly.Treeview")
    files_tree.heading("file", text="Файл")
    files_tree.heading("folder", text="Папка")
    files_tree.column("file", width=220, minwidth=150)
    files_tree.column("folder", width=590, minwidth=200)
    files_tree.grid(row=0, column=0, sticky="nsew")
    scrollbar = ttk.Scrollbar(files_frame, orient="vertical", command=files_tree.yview)
    scrollbar.grid(row=0, column=1, sticky="ns")
    files_tree.configure(yscrollcommand=scrollbar.set)
    file_buttons = ttk.Frame(files_frame, style="Monthly.TFrame")
    file_buttons.grid(row=1, column=0, columnspan=2, sticky="ew", pady=(9, 0))
    file_buttons.columnconfigure(3, weight=1)
    count_var = tk.StringVar(master=window, value="Файлов: 0")
    ttk.Label(file_buttons, textvariable=count_var, style="Monthly.Muted.TLabel").grid(row=0, column=3, sticky="e")

    info_var = tk.StringVar(master=window, value=(
        "ЖК 3, 5, 6, 7 → Планетная. ЖК 13 и 13Н → Петрозаводская.\n"
        "«Тариф» — сумма строки. Пустые и нулевые значения восстанавливаются по ненулевой цене за 1 услугу в том же месяце."
    ))
    info_label = ttk.Label(root_frame, textvariable=info_var, style="Monthly.Muted.TLabel", justify="left", wraplength=960)
    info_label.grid(row=4, column=0, sticky="ew", pady=(0, 9))

    status_var = tk.StringVar(master=window, value="Добавьте один или несколько файлов .xls, .xlsx или .xlsm.")
    status_label = ttk.Label(root_frame, textvariable=status_var, style="Monthly.TLabel", wraplength=960)
    status_label.grid(row=5, column=0, sticky="ew", pady=(0, 7))

    def wrap_labels(event):
        wrap = max(500, event.width - 20)
        info_label.configure(wraplength=wrap)
        status_label.configure(wraplength=wrap)

    root_frame.bind("<Configure>", wrap_labels)
    log_frame = ttk.LabelFrame(root_frame, text="  Сообщения и контроль данных  ", padding=7, style="Monthly.TLabelframe")
    log_frame.grid(row=6, column=0, sticky="nsew")
    log_frame.columnconfigure(0, weight=1)
    log_frame.rowconfigure(0, weight=1)
    log = tk.Text(log_frame, height=6, wrap="word", state="disabled", borderwidth=0, font=("Arial", 10), background="#ffffff", foreground="#2b4151", padx=8, pady=6)
    log.grid(sticky="nsew")
    log_scroll = ttk.Scrollbar(log_frame, orient="vertical", command=log.yview)
    log_scroll.grid(row=0, column=1, sticky="ns")
    log.configure(yscrollcommand=log_scroll.set)
    progressbar = ttk.Progressbar(root_frame, mode="indeterminate")
    progressbar.grid(row=7, column=0, sticky="ew", pady=(10, 10))
    bottom = ttk.Frame(root_frame, style="Monthly.TFrame")
    bottom.grid(row=8, column=0, sticky="ew")
    bottom.columnconfigure(0, weight=1)

    paths = []
    events = queue.Queue()
    cancel_event = threading.Event()
    busy = False
    close_requested = False
    result_path = None
    worker = None
    poll_id = None

    def add_log(text):
        log.configure(state="normal")
        log.insert("end", str(text).rstrip() + "\n")
        if int(log.index("end-1c").split(".")[0]) > 2200:
            log.delete("1.0", "201.0")
        log.see("end")
        log.configure(state="disabled")

    def refresh_files():
        files_tree.delete(*files_tree.get_children())
        for index, path in enumerate(paths):
            files_tree.insert("", "end", iid=str(index), values=(path.name, str(path.parent)))
        count_var.set(f"Файлов: {len(paths)}")
        build_button.configure(state="disabled" if busy or not paths else "normal")
        remove_button.configure(state="disabled" if busy or not paths else "normal")
        clear_button.configure(state="disabled" if busy or not paths else "normal")

    def add_files():
        if busy:
            return
        selected = filedialog.askopenfilenames(
            parent=window, title="Выберите выгрузки за нужные месяцы",
            filetypes=[("Выгрузки Excel", "*.xls *.xlsx *.xlsm"), ("Все файлы", "*.*")],
        )
        known = {os.path.normcase(str(path)) for path in paths}
        rejected = []
        for name in selected:
            path = Path(name).expanduser().resolve()
            if path.suffix.lower() not in SUPPORTED:
                rejected.append(path.name)
                continue
            key = os.path.normcase(str(path))
            if key not in known:
                paths.append(path)
                known.add(key)
        refresh_files()
        if rejected:
            messagebox.showwarning("Формат файла", "Пропущены неподдерживаемые файлы:\n" + "\n".join(rejected), parent=window)
        if selected:
            status_var.set(f"Добавлено файлов: {len(paths)}. Названия файлов не влияют на распределение по месяцам.")

    def remove_files():
        if busy:
            return
        indices = sorted((int(key) for key in files_tree.selection()), reverse=True)
        for index in indices:
            paths.pop(index)
        refresh_files()

    def clear_files():
        if not busy:
            paths.clear()
            refresh_files()

    def set_busy(value):
        nonlocal busy
        busy = value
        add_button.configure(state="disabled" if value else "normal")
        year_input.configure(state="disabled" if value else "normal")
        from_input.configure(state="disabled" if value else "readonly")
        to_input.configure(state="disabled" if value else "readonly")
        cancel_button.configure(state="normal" if value else "disabled")
        open_button.configure(state="disabled" if value or result_path is None else "normal")
        refresh_files()
        if value:
            progressbar.start(12)
        else:
            progressbar.stop()
            progressbar.configure(value=0)

    def run_worker(input_paths, year, first_month, last_month, output, overwrite):
        def progress(message):
            if cancel_event.is_set():
                raise Cancelled("Формирование отчёта отменено.")
            events.put(("progress", str(message)))
        try:
            ensure_dependencies(require_xls=any(path.suffix.lower() == ".xls" for path in input_paths))
            report = build_report(input_paths, year, first_month, last_month, progress=progress, cancel=cancel_event)
            progress("Сохранение Excel-книги…")
            saved = save_report(report, output, input_paths, overwrite=overwrite)
            events.put(("success", (saved, report)))
        except Cancelled as exc:
            events.put(("cancelled", str(exc) or "Формирование отчёта отменено."))
        except (ReportError, OSError, ValueError) as exc:
            events.put(("error", str(exc)))
        except Exception as exc:
            events.put(("error", f"Не удалось сформировать отчёт ({type(exc).__name__}): {exc}"))

    def start_report():
        nonlocal result_path, worker
        if busy or not paths:
            return
        try:
            year = int(year_var.get().strip())
            if not 1900 <= year <= 9999:
                raise ValueError
            first_month = MONTHS.index(start_var.get())
            last_month = MONTHS.index(end_var.get())
            if first_month > last_month:
                messagebox.showerror("Период", "Начальный месяц должен быть раньше конечного или совпадать с ним.", parent=window)
                return
        except ValueError:
            messagebox.showerror("Период", "Введите год числом и выберите месяцы отчёта.", parent=window)
            return
        name = filedialog.asksaveasfilename(
            parent=window, title="Сохранить сводный отчёт", defaultextension=".xlsx",
            initialfile=f"ЖК_Итоги_{year}_{first_month:02d}-{last_month:02d}.xlsx",
            filetypes=[("Excel-книга", "*.xlsx")], confirmoverwrite=False,
        )
        if not name:
            return
        output = Path(name).expanduser().resolve()
        if output.suffix.lower() != ".xlsx":
            messagebox.showerror("Формат результата", "Укажите имя с расширением .xlsx.", parent=window)
            return
        if any(output == source or (output.exists() and source.exists() and os.path.samefile(output, source)) for source in paths):
            messagebox.showerror("Исходный файл", "Выберите другое имя: итоговый отчёт не должен заменять исходную выгрузку.", parent=window)
            return
        overwrite = output.exists()
        if overwrite and not messagebox.askyesno("Заменить файл?", f"Файл уже существует:\n{output}\n\nЗаменить его новым отчётом?", parent=window, default="no"):
            return
        result_path = None
        cancel_event.clear()
        log.configure(state="normal")
        log.delete("1.0", "end")
        log.configure(state="disabled")
        add_log(f"Период: {MONTHS[first_month]} — {MONTHS[last_month]} {year}. Файлов: {len(paths)}.")
        set_busy(True)
        status_var.set("Чтение и проверка исходных данных…")
        worker = threading.Thread(target=run_worker, args=(list(paths), year, first_month, last_month, output, overwrite), daemon=True)
        worker.start()

    def cancel_report():
        if busy:
            cancel_event.set()
            cancel_button.configure(state="disabled")
            status_var.set("Остановка… Дождитесь завершения текущей операции.")

    def open_result():
        if result_path is not None:
            try:
                _open_result_file(result_path)
            except OSError as exc:
                messagebox.showerror("Открыть отчёт", f"Не удалось открыть Excel. Файл сохранён здесь:\n{result_path}\n\n{exc}", parent=window)

    def destroy_window():
        nonlocal poll_id
        if poll_id is not None:
            window.after_cancel(poll_id)
            poll_id = None
        window.destroy()

    def on_close():
        nonlocal close_requested
        if busy:
            if close_requested:
                return
            if not messagebox.askyesno("Отчёт формируется", "Остановить обработку и закрыть окно?", parent=window, default="no"):
                return
            close_requested = True
            cancel_report()
        else:
            destroy_window()

    def poll_events():
        nonlocal result_path, poll_id
        poll_id = None
        for _ in range(120):
            try:
                kind, payload = events.get_nowait()
            except queue.Empty:
                break
            if kind == "progress":
                add_log(payload)
                if not cancel_event.is_set():
                    status_var.set(payload)
                continue
            set_busy(False)
            if kind == "success":
                result_path, report = payload
                open_button.configure(state="normal")
                quantities = sum((value.quantity for value in report.groups.values()), ZERO)
                money = sum((value.money for value in report.groups.values()), ZERO)
                summary = f"Готово. Услуг: {quantities:,.0f}; сумма: {money:,.2f} руб.".replace(",", " ")
                status_var.set(summary)
                add_log(summary)
                add_log(f"Восстановлено пустых и нулевых тарифов: {len(report.restorations)}.")
                for warning in report.warnings:
                    add_log("Внимание: " + str(warning))
                add_log(f"Сохранено: {result_path}")
                if not close_requested:
                    messagebox.showinfo("Отчёт готов", summary + f"\n\n{result_path}" + ("\n\nПроверьте сообщения и лист «Источники»: есть предупреждения." if report.warnings else ""), parent=window)
            elif kind == "cancelled":
                status_var.set("Обработка отменена.")
                add_log(payload)
            else:
                status_var.set("Отчёт не сформирован. Подробности — в сообщениях.")
                add_log("Ошибка: " + payload)
                if not close_requested:
                    # Детальные строки остаются в прокручиваемом журнале окна.
                    # Большой messagebox может выходить за пределы экрана.
                    brief = payload if len(payload) <= 450 else payload.split("\n", 1)[0][:400] + "\n\nПодробности — в поле «Сообщения и контроль данных» основного окна."
                    messagebox.showerror("Не удалось сформировать отчёт", brief, parent=window)
            if close_requested:
                destroy_window()
                return
        poll_id = window.after(100, poll_events)

    add_button = ttk.Button(file_buttons, text="＋ Добавить файлы", command=add_files)
    add_button.grid(row=0, column=0, padx=(0, 8))
    remove_button = ttk.Button(file_buttons, text="Убрать выбранные", command=remove_files)
    remove_button.grid(row=0, column=1, padx=(0, 8))
    clear_button = ttk.Button(file_buttons, text="Очистить", command=clear_files)
    clear_button.grid(row=0, column=2)
    open_button = ttk.Button(bottom, text="Открыть результат", command=open_result, state="disabled")
    open_button.grid(row=0, column=0, sticky="w")
    cancel_button = ttk.Button(bottom, text="Отмена", command=cancel_report, state="disabled")
    cancel_button.grid(row=0, column=1, padx=(10, 10))
    build_button = ttk.Button(bottom, text="Сформировать Excel", command=start_report, style="Monthly.Primary.TButton", state="disabled")
    build_button.grid(row=0, column=2)
    files_tree.bind("<Delete>", lambda _event: remove_files())
    window.protocol("WM_DELETE_WINDOW", on_close)
    refresh_files()
    poll_id = window.after(100, poll_events)
    if parent is None:
        window.mainloop()
    return window


def main(argv=None):
    """CLI: python zhk_monthly_report.py январь.xls февраль.xls --year 2026 --from-month 1 --to-month 10 -o итог.xlsx"""
    parser = argparse.ArgumentParser(
        description="Сводный Excel-отчёт ЖК по месяцам. Без файлов открывается окно.",
        epilog="Пример: python zhk_monthly_report.py январь.xls февраль.xls --year 2026 --from-month 1 --to-month 10 -o Итоги.xlsx",
    )
    parser.add_argument("files", nargs="*", type=Path, help="Исходные выгрузки .xls/.xlsx/.xlsm")
    parser.add_argument("--year", type=int, default=datetime.now().year, help="Год отчёта (по умолчанию текущий)")
    parser.add_argument("--from-month", type=int, choices=range(1, 13), default=1, metavar="1–12", help="Первый месяц (по умолчанию 1)")
    parser.add_argument("--to-month", type=int, choices=range(1, 13), default=datetime.now().month, metavar="1–12", help="Последний месяц (по умолчанию текущий)")
    parser.add_argument("--output", "-o", type=Path, help="Путь итоговой книги .xlsx")
    parser.add_argument("--overwrite", action="store_true", help="Разрешить замену существующего итогового файла")
    args = parser.parse_args(argv)
    if not args.files:
        try:
            launch_monthly()
            return 0
        except ReportError as exc:
            print(f"Ошибка: {exc}", file=sys.stderr)
            return 1
    if args.output is None:
        parser.error("при обработке файлов нужен параметр --output / -o")
    if args.from_month > args.to_month:
        parser.error("--from-month не может быть больше --to-month")
    try:
        ensure_dependencies(require_xls=any(path.suffix.lower() == ".xls" for path in args.files))
        report = build_report(args.files, args.year, args.from_month, args.to_month, progress=lambda text: print(text, flush=True))
        result = save_report(report, args.output, args.files, overwrite=args.overwrite)
        for warning in report.warnings:
            print("Внимание: " + str(warning), file=sys.stderr)
        print(f"Восстановлено пустых и нулевых тарифов: {len(report.restorations)}.")
        print(f"Готово: {result}")
        return 0
    except (KeyboardInterrupt, Cancelled):
        print("Обработка отменена.", file=sys.stderr)
        return 130
    except (ReportError, OSError, ValueError) as exc:
        print(f"Ошибка: {exc}", file=sys.stderr)
        return 1
    except Exception as exc:
        print(f"Не удалось сформировать отчёт ({type(exc).__name__}): {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
